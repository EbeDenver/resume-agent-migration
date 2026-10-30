"""Generate complete application package for BSI (British Standards Institution): AI Technical Specialist.

Job Requisition: JR0020733-1
Organization: BSI (British Standards Institution) — Artificial Intelligence Notified Body / AI Conformity Assessment
Role: AI Technical Specialist
Location: US - Remote / Home-based
Salary Range: $119,400 to $140,000 annually + Comprehensive Benefits

Deliverables:
1. Ebenezer_Rolston_BSI_AI_Technical_Specialist_Resume (DOCX & MD in Bulletted-Standard format)
2. Ebenezer_Rolston_BSI_AI_Technical_Specialist_Skill_Gap_Analysis (DOCX & MD)
3. Ebenezer_Rolston_BSI_AI_Technical_Specialist_Interview_Prep (DOCX & MD)
4. Ebenezer_Rolston_BSI_AI_Technical_Specialist_Cover_Letter (DOCX & MD)
"""
import os
import sys
import shutil
from pathlib import Path

# Add project root src to path
ROOT_DIR = Path(__file__).parent
sys.path.insert(0, str(ROOT_DIR / "src"))

from resume_agent.utils.docx_writer import create_docx

OUTPUT_DIR = ROOT_DIR / "output" / "BSI_AI_Technical_Specialist"
TWIN_DIR = Path("/Users/eberolston/Downloads/AI Agent/ResumeAgent/output/BSI_AI_Technical_Specialist")

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
TWIN_DIR.mkdir(parents=True, exist_ok=True)

def generate_resume_paragraphs():
    p = []

    # 1. Candidate Name
    p.append({
        "align": "center",
        "space_before": 0,
        "space_after": 20,
        "text_runs": [
            {"text": "Ebenezer J. Rolston", "bold": True, "color": "17365D", "size": 36, "font": "Times New Roman"}
        ]
    })

    # Subtitle
    p.append({
        "align": "center",
        "space_before": 0,
        "space_after": 30,
        "text_runs": [
            {"text": "AI Conformity Assessment • Machine Learning Systems • Regulatory Standards & Model Governance", "bold": True, "color": "17365D", "size": 19, "font": "Times New Roman"}
        ]
    })

    # Contact Line
    p.append({
        "align": "center",
        "space_before": 0,
        "space_after": 80,
        "text_runs": [
            {"text": "Denver, CO • ", "size": 18, "color": "1F1F1F"},
            {"text": "eberolston@yahoo.com", "size": 18, "color": "1F1F1F"},
            {"text": " • 501-307-3237 • ", "size": 18, "color": "1F1F1F"},
            {"text": "LinkedIn", "size": 18, "color": "17365D", "bold": True, "url": "https://www.linkedin.com/in/ebenezer-rolston/"},
            {"text": " • ", "size": 18, "color": "1F1F1F"},
            {"text": "GitHub", "size": 18, "color": "17365D", "bold": True, "url": "https://github.com/EbeDenver"},
            {"text": " • ", "size": 18, "color": "1F1F1F"},
            {"text": "AI Guardians", "size": 18, "color": "17365D", "bold": True, "url": "https://portal.aiguardiansllc.us/"}
        ]
    })

    # Section 1: Professional Summary
    p.append({
        "text_runs": [{"text": "PROFESSIONAL SUMMARY", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 120, "space_after": 40
    })

    summary_text = (
        "AI Systems and Conformity Specialist with 20+ years of technical systems engineering, compliance auditing, and "
        "machine learning evaluation experience across regulated government and enterprise environments. Combines deep expertise in "
        "AI regulatory standards (EU AI Act, ISO/IEC 42001, ISO/IEC 23894, NIST AI RMF 1.0) with advanced hands-on data science and machine "
        "learning competencies gained through UC Berkeley (Machine Learning & Deep Learning with Spark - Grade A), Stanford Continuing Studies "
        "(TECH 77 Modern AI Architectures & Python Data Analysis), Harvard SEAS (LLMs & Agentic AI), and Oxford MLx Cases. Experienced in "
        "assessing technical documentation, conducting model validation across supervised/unsupervised learning, neural networks, classification, "
        "and LLMs, and identifying risk vectors in software development lifecycles. Brings 14 years supporting federal platforms at the U.S. "
        "Department of the Interior and 11 years at Acxiom architecting high-volume data validation engines. Holds an M.S. and B.S. in Computer Science "
        "and is a two-time recipient of Department of the Interior Awards for Excellence."
    )
    p.append({
        "text_runs": [{"text": summary_text, "size": 18, "color": "1F1F1F"}],
        "space_before": 0, "space_after": 40
    })

    p.append({
        "text_runs": [{"text": "Winner of two Department of Interior Awards", "italic": True, "color": "5B9BD5", "bold": True, "size": 18}],
        "space_before": 0, "space_after": 80
    })

    # Section 2: Core Competencies (ZERO PARENTHESES POLICY)
    p.append({
        "text_runs": [{"text": "CORE COMPETENCIES", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 100, "space_after": 40
    })

    competencies = (
        "▪ AI Conformity Assessment & Auditing ▪ Machine Learning Model Evaluation ▪ Regulatory Standards & EU AI Act ▪ "
        "ISO 42001 & NIST AI RMF ▪ Deep Learning & Neural Networks ▪ Supervised & Unsupervised Learning ▪ "
        "Distributed ML with Apache Spark ▪ LLM & Multimodal Architecture Assessment ▪ Risk Management & AI Safety ▪ "
        "Software Development Lifecycle SDLC ▪ Technical Documentation & Traceability ▪ Cross-Functional Technical Advisory"
    )
    p.append({
        "text_runs": [{"text": competencies, "size": 17, "color": "1F1F1F", "bold": True}],
        "space_before": 0, "space_after": 100
    })

    # Section 3: Professional Experience
    p.append({
        "text_runs": [{"text": "PROFESSIONAL EXPERIENCE", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 120, "space_after": 40
    })

    # Role 1: AI Guardians LLC
    p.append({
        "text_runs": [
            {"text": "AI Engineering & Governance Lead", "bold": True, "color": "17365D", "size": 19},
            {"text": " | AI Guardians LLC", "color": "1F1F1F", "size": 19}
        ],
        "space_before": 40, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "Denver, CO | Jan 2026–Present", "italic": True, "color": "595959", "size": 17}],
        "space_before": 0, "space_after": 20
    })

    bullets_guardian = [
        "Architected technical assessment frameworks and verification harnesses evaluating machine learning models and generative AI systems for robustness, data drift, and bias.",
        "Operationalized conformity assessment criteria mapped to the EU AI Act High-Risk classifications, ISO/IEC 42001 AI Management Systems, and NIST AI RMF 1.0.",
        "Engineered schema-governed validation layers and Pydantic contracts that enforce deterministic testing gates on model inference, inputs, and outputs.",
        "Built automated testing suites assessing prompt injection resistance, hallucination boundaries, and data lineage traceability across multimodal agent workflows."
    ]
    for b in bullets_guardian:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 20,
            "text_runs": [{"text": b, "size": 18, "color": "1F1F1F"}]
        })

    # Role 2: U.S. Department of the Interior (Program Analyst)
    p.append({
        "text_runs": [
            {"text": "Program Analyst, Acquisition Directorate", "bold": True, "color": "17365D", "size": 19},
            {"text": " | U.S. Department of the Interior", "color": "1F1F1F", "size": 19}
        ],
        "space_before": 60, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "Denver, CO | Jan–Sep 2025", "italic": True, "color": "595959", "size": 17}],
        "space_before": 0, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "Cross-functional Analyst lead for Acquisition Directorate enterprise systems and business technology; GS-13", "italic": True, "color": "17365D", "size": 17}],
        "space_before": 0, "space_after": 20
    })

    bullets_doi = [
        "Senior operational lead for mission-critical enterprise platforms, managing compliance reviews, technical documentation, and high-availability system integrity.",
        "Conducted root-cause technical investigations and data audits across interconnected platforms, assuring adherence to federal regulatory standards and cybersecurity directives.",
        "Authored system compliance documentation, standard operating procedures (SOPs), and audit response artifacts for federal oversight bodies.",
        "Provided technical mentorship and training to operational personnel, bridging complex software engineering requirements with regulatory accountability."
    ]
    for b in bullets_doi:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 20,
            "text_runs": [{"text": b, "size": 18, "color": "1F1F1F"}]
        })

    # Role 3: ARK Solutions / Client: IBC Acquisition Directorate
    p.append({
        "text_runs": [
            {"text": "Senior Technical Systems Consultant", "bold": True, "color": "17365D", "size": 19},
            {"text": " | ARK Solutions — Client : IBC Acquisition Directorate", "color": "1F1F1F", "size": 19}
        ],
        "space_before": 60, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "Denver, CO | 2011–2025", "italic": True, "color": "595959", "size": 17}],
        "space_before": 0, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "14-year technical systems support, architecture validation & compliance review", "italic": True, "color": "17365D", "size": 17}],
        "space_before": 0, "space_after": 20
    })

    bullets_ark = [
        "Served as Senior Escalation Lead handling 5 to 10 mission-critical enterprise incidents daily; performed deep root-cause analysis to eliminate recurring defects.",
        "Department ETL Architect for the multi-agency FBMS conversion, validating data integrity, schema normalization, and system transformations across 4 federal agencies.",
        "Reviewed technical specifications, system architecture documents, and test reports to ensure compliance with federal security and procurement regulations.",
        "Developed automated SQL reconciliation frameworks comparing pre- and post-migration system states to verify complete data fidelity and traceability."
    ]
    for b in bullets_ark:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 20,
            "text_runs": [{"text": b, "size": 18, "color": "1F1F1F"}]
        })

    # Role 4: Acxiom Corporation
    p.append({
        "text_runs": [
            {"text": "Senior Solutions Architect / Systems Analyst", "bold": True, "color": "17365D", "size": 19},
            {"text": " | Acxiom Corporation", "color": "1F1F1F", "size": 19}
        ],
        "space_before": 60, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "Conway, AR | 2000–2011", "italic": True, "color": "595959", "size": 17}],
        "space_before": 0, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "11-year progressive leadership in large-scale data systems, algorithmic scoring & validation", "italic": True, "color": "17365D", "size": 17}],
        "space_before": 0, "space_after": 20
    })

    bullets_acxiom = [
        "Architected enterprise identity matching pipelines, evaluating algorithmic scoring rules and entity resolution models across massive transaction databases.",
        "Conducted rigorous data quality profiling, risk assessments, and model performance audits to ensure compliance with data governance and privacy policies.",
        "Partnered with cross-functional engineering and client teams to design functional test cases, validating model behavior under boundary conditions."
    ]
    for b in bullets_acxiom:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 20,
            "text_runs": [{"text": b, "size": 18, "color": "1F1F1F"}]
        })

    # Section 4: Machine Learning & Data Science Portfolio
    p.append({
        "text_runs": [{"text": "APPLIED MACHINE LEARNING & AI SYSTEMS EVALUATION", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 120, "space_after": 40
    })

    # Berkeley ML Spark
    p.append({
        "text_runs": [
            {"text": "Distributed Machine Learning & Deep Learning with Spark", "bold": True, "color": "17365D", "size": 19},
            {"text": " | UC Berkeley Extension (COMPSCI X433.7, Grade A)", "italic": True, "color": "595959", "size": 17}
        ],
        "space_before": 20, "space_after": 10
    })
    p.append({
        "bullet": True, "space_before": 0, "space_after": 15,
        "text_runs": [{"text": "Engineered and validated distributed machine learning pipelines with PySpark and Spark MLlib, evaluating neural networks, supervised and unsupervised classifiers, regression models, and clustering algorithms across large-scale datasets.", "size": 18, "color": "1F1F1F"}]
    })
    p.append({
        "bullet": True, "space_before": 0, "space_after": 15,
        "text_runs": [{"text": "Conducted hyperparameter tuning, cross-validation, and error surface analysis to measure model variance, bias, and generalizability under rigorous statistical criteria.", "size": 18, "color": "1F1F1F"}]
    })

    # Stanford TECH 77
    p.append({
        "text_runs": [
            {"text": "Auditable Multimodal Document Intelligence & Model Benchmarking", "bold": True, "color": "17365D", "size": 19},
            {"text": " | Stanford TECH 77: Modern AI Architectures", "italic": True, "color": "595959", "size": 17}
        ],
        "space_before": 20, "space_after": 10
    })
    p.append({
        "bullet": True, "space_before": 0, "space_after": 15,
        "text_runs": [{"text": "Benchmarked Vision Language Models (VLMs), Document Transformers, and OCR baselines on unstructured business records, analyzing failure modes, hallucination rates, and confidence-score calibration.", "size": 18, "color": "1F1F1F"}]
    })
    p.append({
        "bullet": True, "space_before": 0, "space_after": 15,
        "text_runs": [{"text": "Architected LangGraph StateGraph conditional routing with strict Pydantic validation and prompt-injection anomaly gating, enforcing verifiable human-in-the-loop escalation paths.", "size": 18, "color": "1F1F1F"}]
    })

    # Oxford MLx Cases
    p.append({
        "text_runs": [
            {"text": "End-to-End AI System Deployment & Robustness Testing", "bold": True, "color": "17365D", "size": 19},
            {"text": " | Oxford Machine Learning Institute (MLx Cases 2026)", "italic": True, "color": "595959", "size": 17}
        ],
        "space_before": 20, "space_after": 10
    })
    p.append({
        "bullet": True, "space_before": 0, "space_after": 20,
        "text_runs": [{"text": "Participated in hands-on system building and assessment modules, evaluating production AI architectures for latency bottlenecks, boundary failure modes, and monitoring instrumentation.", "size": 18, "color": "1F1F1F"}]
    })

    # Section 5: Education & Credentials (DATE-FREE RULE)
    p.append({
        "text_runs": [{"text": "EDUCATION & PROFESSIONAL CREDENTIALS", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 120, "space_after": 40
    })

    edu_items = [
        ("M.S. & B.S. in Computer Science", "St. Joseph's College"),
        ("Machine Learning and Deep Learning with Spark (COMPSCI X433.7, Grade A)", "UC Berkeley Extension"),
        ("TECH 77: Modern AI Architectures & Data Analysis with Python", "Stanford Continuing Studies"),
        ("Graduate Certificate in AI Strategy & Applications in IT", "University of Denver"),
        ("Executive Certificate — LLMs: Transformer to Agentic AI", "Harvard John A. Paulson SEAS"),
        ("AI Governance & Privacy Frameworks (NIST AI RMF, ISO 42001)", "AI Tech and Privacy Academy"),
        ("MLx Cases 2026: Applied AI Systems Building", "Oxford Machine Learning Institute"),
        ("Two Department of the Interior Awards for Excellence", "Interior Business Center & FBMS Program")
    ]
    for degree, school in edu_items:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 15,
            "text_runs": [
                {"text": degree, "bold": True, "color": "17365D", "size": 18},
                {"text": f": {school}", "size": 18, "color": "1F1F1F"}
            ]
        })

    # Section 6: Technical Proficiencies
    p.append({
        "text_runs": [{"text": "TECHNICAL & CONFORMITY PROFICIENCIES", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 120, "space_after": 40
    })

    proficiencies = [
        ("AI Regulations, Standards & Notified Body Criteria", "EU AI Act High-Risk Requirements • ISO/IEC 42001 AI Management Systems • ISO/IEC 23894 Risk Management • NIST AI RMF 1.0 • CE Marking Technical Documentation • ISO 13485 & MDR/IVDR AI Device Principles • Conformity Assessment Procedures"),
        ("Machine Learning & Deep Learning Architectures", "Supervised & Unsupervised Learning • Neural Networks • Deep Learning • Classification & Regression • Clustering • PySpark & Spark MLlib • Large Language Models (LLMs) • Vision Language Models (VLMs) • Document Transformers • Probabilistic Methods"),
        ("Model Evaluation, Testing & Validation", "Model Performance Metrics (ROC-AUC, Precision, Recall, F1) • Cross-Validation & Hyperparameter Optimization • Data Drift & Distribution Shifts • Algorithmic Bias & Fairness Auditing • Adversarial & Prompt Injection Testing • Hallucination Gating"),
        ("Software Engineering, SDLC & Governance", "Technical Documentation Review • Traceability Matrices • Risk Management & Mitigation • Python (Pandas, NumPy, Scikit-Learn) • Pydantic Contract Enforcement • REST API & JSON Schema • Git & GitHub • Stored Procedures & Relational SQL"),
        ("Technical Advisory & Stakeholder Communication", "Conformity Assessment Technical Reporting • Client Technical Consultation • Internal Team Mentoring • Cross-Functional Consensus Building • Standard Operating Procedures (SOPs) • Auditable Work Product Authoring")
    ]
    for category, skills in proficiencies:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 15,
            "text_runs": [
                {"text": f"{category}: ", "bold": True, "color": "17365D", "size": 18},
                {"text": skills, "size": 18, "color": "1F1F1F"}
            ]
        })

    return p

def generate_resume_md():
    return """# Ebenezer J. Rolston
**AI Conformity Assessment • Machine Learning Systems • Regulatory Standards & Model Governance**  
Denver, CO • eberolston@yahoo.com • 501-307-3237 • [LinkedIn](https://www.linkedin.com/in/ebenezer-rolston/) • [GitHub](https://github.com/EbeDenver) • [AI Guardians](https://portal.aiguardiansllc.us/)

---

## PROFESSIONAL SUMMARY

AI Systems and Conformity Specialist with 20+ years of technical systems engineering, compliance auditing, and machine learning evaluation experience across regulated government and enterprise environments. Combines deep expertise in AI regulatory standards (EU AI Act, ISO/IEC 42001, ISO/IEC 23894, NIST AI RMF 1.0) with advanced hands-on data science and machine learning competencies gained through UC Berkeley (Machine Learning & Deep Learning with Spark - Grade A), Stanford Continuing Studies (TECH 77 Modern AI Architectures & Python Data Analysis), Harvard SEAS (LLMs & Agentic AI), and Oxford MLx Cases. Experienced in assessing technical documentation, conducting model validation across supervised/unsupervised learning, neural networks, classification, and LLMs, and identifying risk vectors in software development lifecycles. Brings 14 years supporting federal platforms at the U.S. Department of the Interior and 11 years at Acxiom architecting high-volume data validation engines. Holds an M.S. and B.S. in Computer Science and is a two-time recipient of Department of the Interior Awards for Excellence.

*Winner of two Department of Interior Awards*

---

## CORE COMPETENCIES

▪ AI Conformity Assessment & Auditing ▪ Machine Learning Model Evaluation ▪ Regulatory Standards & EU AI Act ▪ ISO 42001 & NIST AI RMF ▪ Deep Learning & Neural Networks ▪ Supervised & Unsupervised Learning ▪ Distributed ML with Apache Spark ▪ LLM & Multimodal Architecture Assessment ▪ Risk Management & AI Safety ▪ Software Development Lifecycle SDLC ▪ Technical Documentation & Traceability ▪ Cross-Functional Technical Advisory

---

## PROFESSIONAL EXPERIENCE

### AI Engineering & Governance Lead | AI Guardians LLC
*Denver, CO | Jan 2026–Present*
- Architected technical assessment frameworks and verification harnesses evaluating machine learning models and generative AI systems for robustness, data drift, and bias.
- Operationalized conformity assessment criteria mapped to the EU AI Act High-Risk classifications, ISO/IEC 42001 AI Management Systems, and NIST AI RMF 1.0.
- Engineered schema-governed validation layers and Pydantic contracts that enforce deterministic testing gates on model inference, inputs, and outputs.
- Built automated testing suites assessing prompt injection resistance, hallucination boundaries, and data lineage traceability across multimodal agent workflows.

### Program Analyst, Acquisition Directorate | U.S. Department of the Interior
*Denver, CO | Jan–Sep 2025*  
*Cross-functional Analyst lead for Acquisition Directorate enterprise systems and business technology; GS-13*
- Senior operational lead for mission-critical enterprise platforms, managing compliance reviews, technical documentation, and high-availability system integrity.
- Conducted root-cause technical investigations and data audits across interconnected platforms, assuring adherence to federal regulatory standards and cybersecurity directives.
- Authored system compliance documentation, standard operating procedures (SOPs), and audit response artifacts for federal oversight bodies.
- Provided technical mentorship and training to operational personnel, bridging complex software engineering requirements with regulatory accountability.

### Senior Technical Systems Consultant | ARK Solutions — Client : IBC Acquisition Directorate
*Denver, CO | 2011–2025*  
*14-year technical systems support, architecture validation & compliance review*
- Served as Senior Escalation Lead handling 5 to 10 mission-critical enterprise incidents daily; performed deep root-cause analysis to eliminate recurring defects.
- Department ETL Architect for the multi-agency FBMS conversion, validating data integrity, schema normalization, and system transformations across 4 federal agencies.
- Reviewed technical specifications, system architecture documents, and test reports to ensure compliance with federal security and procurement regulations.
- Developed automated SQL reconciliation frameworks comparing pre- and post-migration system states to verify complete data fidelity and traceability.

### Senior Solutions Architect / Systems Analyst | Acxiom Corporation
*Conway, AR | 2000–2011*  
*11-year progressive leadership in large-scale data systems, algorithmic scoring & validation*
- Architected enterprise identity matching pipelines, evaluating algorithmic scoring rules and entity resolution models across massive transaction databases.
- Conducted rigorous data quality profiling, risk assessments, and model performance audits to ensure compliance with data governance and privacy policies.
- Partnered with cross-functional engineering and client teams to design functional test cases, validating model behavior under boundary conditions.

---

## APPLIED MACHINE LEARNING & AI SYSTEMS EVALUATION

### Distributed Machine Learning & Deep Learning with Spark | UC Berkeley Extension (COMPSCI X433.7, Grade A)
- Engineered and validated distributed machine learning pipelines with PySpark and Spark MLlib, evaluating neural networks, supervised and unsupervised classifiers, regression models, and clustering algorithms across large-scale datasets.
- Conducted hyperparameter tuning, cross-validation, and error surface analysis to measure model variance, bias, and generalizability under rigorous statistical criteria.

### Auditable Multimodal Document Intelligence & Model Benchmarking | Stanford TECH 77: Modern AI Architectures
- Benchmarked Vision Language Models (VLMs), Document Transformers, and OCR baselines on unstructured business records, analyzing failure modes, hallucination rates, and confidence-score calibration.
- Architected LangGraph StateGraph conditional routing with strict Pydantic validation and prompt-injection anomaly gating, enforcing verifiable human-in-the-loop escalation paths.

### End-to-End AI System Deployment & Robustness Testing | Oxford Machine Learning Institute (MLx Cases 2026)
- Participated in hands-on system building and assessment modules, evaluating production AI architectures for latency bottlenecks, boundary failure modes, and monitoring instrumentation.

---

## EDUCATION & PROFESSIONAL CREDENTIALS

- **M.S. & B.S. in Computer Science:** St. Joseph's College
- **Machine Learning and Deep Learning with Spark (COMPSCI X433.7, Grade A):** UC Berkeley Extension
- **TECH 77: Modern AI Architectures & Data Analysis with Python:** Stanford Continuing Studies
- **Graduate Certificate in AI Strategy & Applications in IT:** University of Denver
- **Executive Certificate — LLMs: Transformer to Agentic AI:** Harvard John A. Paulson SEAS
- **AI Governance & Privacy Frameworks (NIST AI RMF, ISO 42001):** AI Tech and Privacy Academy
- **MLx Cases 2026: Applied AI Systems Building:** Oxford Machine Learning Institute
- **Two Department of the Interior Awards for Excellence:** Interior Business Center & FBMS Program

---

## TECHNICAL & CONFORMITY PROFICIENCIES

- **AI Regulations, Standards & Notified Body Criteria:** EU AI Act High-Risk Requirements • ISO/IEC 42001 AI Management Systems • ISO/IEC 23894 Risk Management • NIST AI RMF 1.0 • CE Marking Technical Documentation • ISO 13485 & MDR/IVDR AI Device Principles • Conformity Assessment Procedures
- **Machine Learning & Deep Learning Architectures:** Supervised & Unsupervised Learning • Neural Networks • Deep Learning • Classification & Regression • Clustering • PySpark & Spark MLlib • Large Language Models (LLMs) • Vision Language Models (VLMs) • Document Transformers • Probabilistic Methods
- **Model Evaluation, Testing & Validation:** Model Performance Metrics (ROC-AUC, Precision, Recall, F1) • Cross-Validation & Hyperparameter Optimization • Data Drift & Distribution Shifts • Algorithmic Bias & Fairness Auditing • Adversarial & Prompt Injection Testing • Hallucination Gating
- **Software Engineering, SDLC & Governance:** Technical Documentation Review • Traceability Matrices • Risk Management & Mitigation • Python (Pandas, NumPy, Scikit-Learn) • Pydantic Contract Enforcement • REST API & JSON Schema • Git & GitHub • Stored Procedures & Relational SQL
- **Technical Advisory & Stakeholder Communication:** Conformity Assessment Technical Reporting • Client Technical Consultation • Internal Team Mentoring • Cross-Functional Consensus Building • Standard Operating Procedures (SOPs) • Auditable Work Product Authoring
"""

def generate_skill_gap_md():
    return """# BSI Skill Gap & Fit Analysis
**AI Technical Specialist | AI Notified Body / Conformity Assessment | Requisition JR0020733-1**  
Candidate: Ebenezer J. Rolston (Denver, CO) • Target Role: US Remote / Home-Based • Overall Fit Score: 95/100 • Estimated Interview Probability: 91%

---

## 1. Executive Summary & Strategic Context

The British Standards Institution (BSI) is a world-renowned national standards body and accredited Notified Body. In response to global AI regulatory developments—most notably the **EU AI Act**, **ISO/IEC 42001**, and evolving medical device regulations (MDR/IVDR)—BSI is aggressively expanding its **Artificial Intelligence Notified Body** and AI Conformity Assessment team. 

The **AI Technical Specialist** (Requisition JR0020733-1) plays a pivotal role in reviewing manufacturers' and developers' technical documentation, evaluating AI/ML models across industries (medical devices, robotics, biometrics, financial systems), conducting conformity assessments for CE marking, and shaping internal testing procedures, tools, and training.

Ebenezer J. Rolston offers a compelling synthesis of **mature systems auditing and compliance rigor** (20+ years of enterprise engineering, 14 years supporting federal systems at the U.S. Department of the Interior, two DOI Awards) and **rigorous, hands-on data science and machine learning capabilities**. By grounding his technical evaluation skills in **UC Berkeley's Machine Learning and Deep Learning with Spark (Grade A)**, **Stanford Continuing Studies TECH 77 (Modern AI Architectures)**, **Harvard SEAS Executive Certificate (LLMs & Agentic AI)**, and formal credentials in **AI Governance & Standards (ISO 42001, NIST AI RMF, DU AI Strategy)**, he meets both the regulatory auditing expectations and deep technical model-evaluation criteria of BSI's Notified Body.

---

## 2. Quantitative Weighted Fit Matrix

| Category | Weight | Score | Evaluation Rationale & Evidence |
| :--- | :---: | :---: | :--- |
| **Must-Have Core Skills** | 40% | 38 / 40 | Bachelor's degree equivalent in CS/Software Engineering (verified: M.S. and B.S. in CS); 4+ years post-graduate experience in technical systems / AI model evaluation (20+ yrs enterprise systems, 4+ yrs applied AI/ML study & governance); review of technical documentation; understanding of SDLC and risk management. |
| **Applied AI/ML Knowledge (2+ Areas)** | 25% | 25 / 25 | Exceeds the two-area requirement: Deep learning & neural networks, supervised & unsupervised learning, classification & regression (UC Berkeley COMPSCI X433.7 Grade A); Vision Language Models, Transformers, and LLMs (Stanford TECH 77, Harvard SEAS); distributed data pipelines with Spark MLlib. |
| **AI Regulations & Standards** | 15% | 15 / 15 | Formal certification in AI Governance (AI Tech and Privacy Academy); Graduate Certificate in AI Strategy & Credly badge in AI Ethics (University of Denver); deep operational mastery of EU AI Act, ISO/IEC 42001, ISO/IEC 23894, and NIST AI RMF 1.0. |
| **Education & Credentials** | 10% | 10 / 10 | M.S. and B.S. in Computer Science; UC Berkeley Extension ML & DL with Spark (Grade A); Stanford TECH 77 Modern AI Architectures; Harvard SEAS Executive Certificate; Oxford MLx Cases. |
| **Communication & Global Remote Culture** | 10% | 7 / 10 | 100% remote self-direction; 14 years training 85+ enterprise users and writing SOPs; technical advisory experience across technical and executive stakeholders. (Minor bridge: medical device specific CE-marking audits). |
| **Total Weighted Score** | **100%** | **95 / 100** | **Band: Strong Match (Interview Probability: 91%)** |

---

## 3. Requirement-by-Requirement Evidence Mapping

### 1. Review Technical Documentation & Assess AI Systems Against Standards — [MATCHED (STRONG)]
* **BSI Requirement:** Review technical documentation and assess AI systems against applicable regulatory requirements and standards.
* **Evidence:** Over 14 years at the Department of the Interior conducting formal compliance reviews, technical system evaluations, and Inspector General audit defenses. At AI Guardians, authored technical evaluation checklists and verification scorecards mapped directly to the EU AI Act High-Risk requirements, ISO/IEC 42001, and NIST AI RMF 1.0.

### 2. Experience in at Least Two AI/ML Areas (Deep Learning, Neural Networks, Supervised/Unsupervised, LLMs) — [MATCHED (EXCEEDS)]
* **BSI Requirement:** Experience in at least two AI/ML areas, such as deep learning, neural networks, supervised/unsupervised learning, classification, regression, fuzzy logic, probabilistic methods, optimization, LLMs, or biometrics.
* **Evidence:**
  - **UC Berkeley Extension (COMPSCI X433.7 - Grade A):** Built and evaluated distributed machine learning pipelines with PySpark and Spark MLlib; implemented neural networks, deep learning models, random forests, gradient boosting, logistic regression, k-means clustering, and PCA dimensionality reduction.
  - **Stanford TECH 77 (Modern AI Architectures):** Benchmarked Vision Language Models (VLMs), Document Transformers, and OCR architectures; implemented LangGraph StateGraph conditional routing, Pydantic schema validation, and prompt-injection anomaly gating.
  - **Harvard SEAS Executive Certificate:** Architected and evaluated Large Language Models (LLMs) and agentic multi-agent systems.

### 3. Knowledge of AI Regulations, Standards, and Relevant Guidance — [MATCHED (EXCEEDS)]
* **BSI Requirement:** Knowledge of AI regulations, standards, and relevant guidance (EU AI Act, ISO/IEC 42001, ISO/IEC 23894, NIST AI RMF, ISO 13485/MDR).
* **Evidence:** Formal certification in AI Governance (AI Tech and Privacy Academy, Jan 2026), University of Denver Graduate Certificate in AI Strategy & Credly badge in AI Literacy, Ethics, and Governance (June 2026). Published comprehensive governance blueprints assessing generative AI risk across data ingestion, model validation, bias mitigation, and post-market surveillance.

### 4. Software Development Lifecycles (SDLC) and Risk Management — [MATCHED (EXCEEDS)]
* **BSI Requirement:** Knowledge of software development lifecycles and risk management.
* **Evidence:** 20+ years of SDLC experience across Agile, Waterfall, and hybrid frameworks. Managed enterprise release gates, risk registers, and defect triage as Senior Escalation Lead at DOI/IBC. Implemented automated regression testing and schema validation gates at AI Guardians.

### 5. AI Technical Expertise, Mentoring, and Client Guidance — [MATCHED (EXCEEDS)]
* **BSI Requirement:** Provide AI technical expertise, mentoring, and guidance to colleagues; participate in client meetings and support certification activities.
* **Evidence:** Trained 85+ professionals in technical workflows and system operations; authored comprehensive Standard Operating Procedures (SOPs); regularly conducts technical walkthroughs and advisory sessions translating complex machine learning concepts into actionable compliance milestones.

---

## 4. Gap Identification & Strategic Mitigation Blueprints

| Perceived Gap | BSI Notified Body Perspective | Strategic Mitigation & Positioning Blueprint |
| :--- | :--- | :--- |
| **Direct Notified Body / CE Marking Audit Experience** | BSI may look for prior experience working inside an EU Notified Body or writing formal CE marking certificates for medical devices. | **Positioning Strategy:** Reframe 14 years of federal agency compliance audits (DOI/IBC) under strict statutory frameworks (FAR, FISMA, NIST) as directly equivalent in audit discipline. Emphasize that conformity assessment against standards like ISO/IEC 42001 and EU AI Act requires the exact same evidence-based verification, traceability matrices, and audit defense methodology candidate has practiced for decades. |
| **Regulated Medical Device (MDR/IVDR) Specificity** | Posting mentions medical devices as one of several industry verticals (alongside robotics and biometrics). | **Positioning Strategy:** Emphasize core ML model verification principles that apply universally to high-risk AI: dataset quality, validation/test splits, overfitting, distribution shifts, drift detection, and confidence calibration. Highlight understanding of ISO 13485 software lifecycle principles (IEC 62304) and risk management (ISO 14971) as natural extensions of ISO 42001 and NIST AI RMF. |
| **Practitioner vs. Academic AI/ML Split** | 20+ years in enterprise data and systems, with formal data science and ML training through Berkeley, Stanford, and Harvard. | **Positioning Strategy:** Position this as the ideal profile for an AI Technical Specialist: mature software engineering instincts prevent being misled by vendor claims, while cutting-edge training in Spark MLlib, neural networks, and LLMs enables rigorous technical interrogation of manufacturers' model documentation. |

---

## 5. Recruiter & Search Committee Positioning Script

* **30-Second Elevator Pitch:**  
  *"Ebenezer Rolston brings a rare and powerful combination to BSI’s AI Notified Body: over 20 years of rigorous enterprise software compliance and systems auditing at the U.S. Department of the Interior (two DOI Awards), paired with formal, hands-on machine learning and deep learning mastery from UC Berkeley (COMPSCI X433.7 Grade A in ML & Deep Learning with Spark), Stanford (TECH 77 Modern AI Architectures), and Harvard SEAS. Certified in AI Governance and ISO 42001 / NIST AI RMF frameworks, he possesses the technical depth to evaluate complex neural networks, LLMs, and distributed ML pipelines, alongside the auditing rigor required to verify technical documentation for EU AI Act conformity assessment."*
"""

def generate_interview_prep_md():
    return """# BSI Interview Preparation Package
**AI Technical Specialist | Artificial Intelligence Notified Body | Requisition JR0020733-1**  
Candidate: Ebenezer J. Rolston (Denver, CO) • Prepared for BSI Conformity Assessment Leadership & Technical Interview Panel

---

## 1. BSI & AI Notified Body Strategic Context

The British Standards Institution (BSI) was incorporated by Royal Charter and is the world's first national standards organization. As an accredited Notified Body in Europe (BSI Group The Netherlands B.V., NB 2797), BSI is positioned at the forefront of global AI regulation. Under the **EU AI Act**, providers of High-Risk AI systems (e.g. AI-enabled medical devices, critical infrastructure, biometrics) must undergo rigorous **Conformity Assessments** before placing systems on the EU market and bearing the CE mark.

### Key Dimensions of the AI Technical Specialist Role:
- **Technical Documentation Assessment:** Scrutinizing the manufacturer's technical dossier—including training data provenance, data labeling protocols, model architecture, validation methodologies, test metrics, and cybersecurity safeguards—against harmonized standards (e.g., ISO/IEC 42001, ISO/IEC 23894, ISO/IEC 22989).
- **Evaluating Model Robustness & Generalizability:** Interrogating whether the model has been adequately tested against distribution shifts, adversarial attacks, edge-case failure modes, and demographic bias.
- **Auditing the AI Quality Management System (QMS):** Ensuring the developer maintains continuous risk management, post-market monitoring, incident reporting, and change management procedures throughout the AI lifecycle.
- **Client Technical Engagement & Advisory:** Engaging with client engineering and regulatory affairs teams in technical meetings, communicating non-conformities with precision, and guiding them on certification expectations.

---

## 2. 30-60-90 Day Strategic Onboarding Charter

### Days 1–30: Master BSI Assessment Procedures & Notified Body Protocols
- **Internalize BSI AI SOPs:** Review BSI’s accredited conformity assessment workflows, technical documentation review checklists, and CE-marking reporting templates.
- **Deep-Dive Regulatory Harmonization:** Map EU AI Act Article 9 (Risk Management System), Article 10 (Data and Data Governance), Article 11 (Technical Documentation), and Article 14 (Human Oversight) to BSI assessment tools.
- **Shadow Senior Assessors:** Shadow active conformity assessment reviews across medical device AI software and enterprise AI applications to observe interview and auditing techniques.
- **Audit Tooling Setup:** Familiarize with BSI internal platforms, documentation repositories, and collaborative assessment tracking tools.

### Days 31–60: Conduct Co-Assessments & Interrogate Technical Dossiers
- **Lead Technical Documentation Reviews:** Under mentor supervision, independently review manufacturers' technical files, verifying algorithmic choices, validation metrics (ROC-AUC, F1, calibration curves), and test data segregation.
- **Participate in Client Conformity Meetings:** Attend technical discovery sessions with manufacturers, formulating targeted questions regarding data preprocessing, validation splits, and model drift mitigation.
- **Draft Formal Technical Assessment Reports:** Author clear, defensible review reports detailing compliance findings, observations, and structured requests for additional information (RFIs).
- **Contribute to Internal Tooling:** Collaborate on developing or validating internal AI testing tools used to evaluate client model submissions.

### Days 61–90: Full Assessment Ownership & Team Mentorship
- **Independent Conformity Assessment Sign-Off:** Execute end-to-end technical documentation reviews for assigned AI systems, issuing formal assessment recommendations for CE marking.
- **Develop Specialized Assessment Guidance:** Author technical guidance or training modules for BSI colleagues on evaluating modern generative AI, LLM fine-tuning, and multimodal model risks (leveraging Stanford TECH 77 and Harvard agentic insights).
- **Drive Continuous Improvement:** Identify optimization opportunities in BSI’s technical review workflows, reducing review cycle times while maintaining uncompromised auditing rigor.

---

## 3. High-Probability Technical & Behavioral Scenarios (STAR Format)

### Scenario 1: Evaluating Technical Documentation for Model Robustness and Data Governance
* **Context:** A manufacturer submits an AI diagnostic model claiming 98% accuracy. As an AI Technical Specialist, you must evaluate the validity of their technical file under conformity assessment requirements.
* **Situation:** In my AI evaluation work and academic projects (UC Berkeley ML with Spark and Stanford TECH 77), I have frequently encountered models that perform exceptionally well in laboratory benchmarks but fail under real-world data shifts.
* **Task:** Establish an evidence-based audit protocol to verify whether a model's reported accuracy reflects genuine generalization or subtle data leakage.
* **Action:** I scrutinize the training, validation, and test dataset splits to verify strict temporal or patient/subject isolation, ensuring zero feature contamination. I inspect the data provenance, labeling protocols, and demographic stratification. I require the manufacturer to provide confusion matrices, subgroup performance metrics, calibration curves, and error surface analysis under noise and adversarial perturbations.
* **Result:** Successfully detects hidden overfitting and dataset bias before certification, protecting end users and upholding the integrity of the CE mark.

### Scenario 2: Interrogating Distributed ML Pipelines (UC Berkeley ML with Spark)
* **Context:** An enterprise submits an AI system leveraging large-scale distributed data processing and neural networks. BSI requires deep technical fluency in distributed ML architectures.
* **Situation:** During my UC Berkeley coursework (COMPSCI X433.7 - Grade A), I developed distributed ML pipelines using PySpark and Spark MLlib, evaluating how parallelized training and distributed feature extraction impact model convergence.
* **Task:** Assess whether a distributed ML pipeline maintains numerical stability, reproducibility, and deterministic data handling during scale-out training.
* **Action:** I inspect the data partitioning strategies, feature transformation pipelines, and gradient aggregation mechanisms. I verify that random seeds are fixed and that distributed join operations do not introduce non-deterministic ordering or duplicate records. I review hyperparameter optimization logs (grid search / Bayesian optimization) to verify that cross-validation was performed correctly across Spark worker nodes without test-set leakage.
* **Result:** Demonstrates deep peer-level engineering credibility, validating that complex distributed ML systems satisfy conformity assessment standards for reproducibility and auditability.

### Scenario 3: Auditing LLM & Multimodal Systems for Hallucination and Prompt Injection (Stanford TECH 77)
* **Context:** BSI is assessing generative AI and multimodal applications where stochastic behavior poses regulatory risks under EU AI Act transparency and risk management mandates.
* **Situation:** In my Stanford TECH 77 project, I built and evaluated an auditable document intelligence pipeline benchmarking Vision Language Models (VLMs) and Document Transformers against OCR baselines.
* **Task:** Design verification criteria to ensure generative models operating in regulated domains do not hallucinate critical outputs or fall prey to prompt injection attacks.
* **Action:** I evaluated failure modes where VLMs hallucinatory tokens on complex table structures. To mitigate this, I implemented LangGraph StateGraph deterministic routing, strict Pydantic JSON schema enforcement, and confidence-score anomaly gating that triggered human-in-the-loop review. I also introduced adversarial input filters to catch indirect prompt injection.
* **Result:** Established a concrete, auditable framework for certifying that generative AI systems have adequate deterministic boundaries, output verification gates, and human oversight mechanisms as mandated by Article 14 of the EU AI Act.

### Scenario 4: Communicating Non-Conformities to Client Leadership Without Creating Impasse
* **Context:** During an audit, an AI developer pushes back strongly against a finding of non-conformity regarding inadequate data governance or lack of post-market drift monitoring.
* **Situation:** As a Senior Escalation Lead at DOI/IBC for 14 years, I regularly navigated high-stakes technical disagreements with senior government stakeholders and vendor software engineering leads.
* **Task:** Maintain uncompromised compliance standards while preserving a constructive, professional partnership with the client.
* **Action:** I anchor the discussion strictly in objective, published standard clauses (e.g., ISO/IEC 42001 Section 8 or EU AI Act Article 10). Rather than criticizing the design, I present empirical evidence of the risk—such as lack of performance monitoring under out-of-distribution inputs. I clearly explain the regulatory rationale and outline acceptable remediation pathways (e.g., establishing automated drift alerts and retraining triggers).
* **Result:** Defused client defensiveness, achieved full agreement on corrective actions, and ensured the system met all regulatory criteria prior to formal sign-off.

### Scenario 5: Reviewing Software Development Lifecycles (SDLC) and Change Management
* **Context:** AI systems are dynamic; continuous model updates and retrained weights can invalidate previous conformity assessments.
* **Situation:** Manufacturers often push continuous integration updates to AI models without evaluating whether the change constitutes a "substantial modification" under regulatory frameworks.
* **Task:** Evaluate the manufacturer's change management procedure and SDLC controls to ensure post-certification updates remain compliant.
* **Action:** I audit their version control, model registry, and CI/CD pipelines. I verify that automated regression test suites run on every commit and that significant updates—such as retraining on new demographic distributions or changing model architectures—trigger formal impact assessments and re-notification protocols.
* **Result:** Guarantees that the manufacturer has robust lifecycle controls preventing unmonitored model drift or unauthorized functional alterations in production.

---

## 4. Strategic Reverse Questions for the BSI Interview Panel

### Questions for the Head of AI Conformity Assessment / Notified Body Leadership:
1. *"With the phased enforcement of the EU AI Act progressing rapidly, what are the primary industry sectors currently seeking Notified Body conformity assessments through BSI (e.g., medical devices vs. industrial automation vs. financial biometrics)?"*
2. *"How does BSI coordinate between its Netherlands Notified Body (NB 2797) and the global team of AI Technical Specialists to ensure consistent interpretation of harmonized standards?"*
3. *"What internal testing tools or sandboxes is BSI currently developing to assist assessors in validating client model technical files?"*

### Questions for Senior Technical Assessors / Team Leads:
4. *"When evaluating proprietary or closed-source foundation models embedded in client applications, how does BSI verify compliance with transparency and data provenance requirements?"*
5. *"How do you balance assessing purely statistical ML models (like Spark MLlib classification or regression) versus modern generative LLM / agentic architectures?"*
6. *"What does the typical lifecycle of a technical documentation review look like from initial dossier receipt to final assessment report?"*

### Questions on Professional Development & Team Collaboration:
7. *"Given BSI’s Royal Charter commitment to inspiring trust, how does the AI team contribute to shaping international standards committees (such as ISO/IEC JTC 1/SC 42 or CEN-CENELEC JTC 21)?"*
8. *"What does success look like for an AI Technical Specialist in this team over their first 6 to 12 months?"*
"""

def generate_cover_letter_md():
    return """# Ebenezer J. Rolston
**Denver, CO • eberolston@yahoo.com • 501-307-3237 • [LinkedIn](https://www.linkedin.com/in/ebenezer-rolston/) • [GitHub](https://github.com/EbeDenver)**

October 29, 2026

AI Conformity Assessment Search Committee  
BSI (British Standards Institution)  
US - Remote / Home-Based  

**Subject: Application for AI Technical Specialist (Requisition #JR0020733-1)**

Dear Members of the Search Committee,

I am writing to express my enthusiastic interest in the **AI Technical Specialist** position (Requisition #JR0020733-1) within BSI’s expanding Artificial Intelligence Notified Body. As a standards-driven technologist with over 20 years of enterprise systems engineering, regulatory compliance, and auditing experience, coupled with formal graduate and applied training in machine learning and AI governance, I am inspired by BSI’s mission to inspire trust in digital transformation and AI-enabled systems.

Throughout my career, I have specialized in the rigorous verification of complex, mission-critical software systems operating under strict statutory oversight. Supporting the **U.S. Department of the Interior (DOI) Acquisition Directorate and IBC** for 14 years as a cross-functional Analyst Lead and Senior Escalation Lead, I managed technical system audits, authored standard operating procedures (SOPs), and defended enterprise data integrity through formal federal Inspector General oversight—work recognized with two Department of the Interior Awards for Excellence. This background instilled in me an uncompromising auditing discipline, deep familiarity with traceability matrices, and the ability to review complex technical documentation for regulatory fidelity.

To complement my systems compliance foundation, I have dedicated extensive effort to mastering modern data science, machine learning, and AI regulatory frameworks:
- **Applied Machine Learning & Deep Learning with Spark (UC Berkeley Extension, COMPSCI X433.7 - Grade A):** Engineered and evaluated distributed machine learning pipelines using PySpark and Spark MLlib, conducting cross-validation, hyperparameter tuning, and performance evaluation across neural networks, supervised classifiers, regression models, and clustering algorithms.
- **Multimodal Architectures & Document Intelligence (Stanford TECH 77):** Benchmarked Vision Language Models (VLMs), Document Transformers, and OCR baselines on unstructured business records, implementing LangGraph StateGraph conditional routing with deterministic Pydantic schema validation and prompt-injection anomaly gating to ensure auditable, predictable outputs.
- **AI Regulatory Standards & Governance:** Formally certified in AI Governance (AI Tech and Privacy Academy), holding a Graduate Certificate in AI Strategy and Credly badge in AI Ethics from the University of Denver, and an Executive Certificate in LLMs and Agentic AI from Harvard SEAS. I have authored comprehensive governance frameworks operationalizing the **EU AI Act High-Risk mandates**, **ISO/IEC 42001**, **ISO/IEC 23894**, and **NIST AI RMF 1.0**.

The AI Technical Specialist role requires a unique balance: the technical depth to interrogate complex algorithmic choices, neural network architectures, and training data provenance, alongside the structured communication skills required to articulate non-conformities clearly to client executives and internal assessment teams. Having trained over 85 professionals and led cross-functional technical teams throughout my career, I excel at translating complex technical requirements into defensible, high-integrity audit outcomes.

As an autonomous, self-directed professional operating from the Denver metropolitan area, I welcome the opportunity to support BSI’s global AI Notified Body in shaping the future of AI conformity assessment. Thank you for your consideration, and I look forward to discussing how my experience can contribute to BSI’s distinguished legacy of excellence.

Sincerely,

**Ebenezer J. Rolston**  
Denver, CO | 501-307-3237 | eberolston@yahoo.com
"""

def generate_cover_letter_docx(docx_path):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    run_name = title.add_run("EBENEZER J. ROLSTON\n")
    run_name.font.name = "Calibri"
    run_name.font.size = Pt(16)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(15, 34, 64)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(14)
    run_sub = sub.add_run("Denver, CO • eberolston@yahoo.com • 501-307-3237 • linkedin.com/in/ebenezer-rolston • github.com/EbeDenver")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(9.5)
    run_sub.font.color.rgb = RGBColor(90, 105, 120)

    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_after = Pt(10)
    r = p_date.add_run("October 29, 2026\n\nAI Conformity Assessment Search Committee\nBSI (British Standards Institution)\nUS - Remote / Home-Based")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(40, 50, 60)

    p_subj = doc.add_paragraph()
    p_subj.paragraph_format.space_after = Pt(10)
    r_subj = p_subj.add_run("Subject: Application for AI Technical Specialist (Requisition #JR0020733-1)")
    r_subj.font.name = "Calibri"
    r_subj.font.size = Pt(10.5)
    r_subj.font.bold = True
    r_subj.font.color.rgb = RGBColor(15, 34, 64)

    paragraphs = [
        "Dear Members of the Search Committee,",
        "I am writing to express my enthusiastic interest in the AI Technical Specialist position (Requisition #JR0020733-1) within BSI’s expanding Artificial Intelligence Notified Body. As a standards-driven technologist with over 20 years of enterprise systems engineering, regulatory compliance, and auditing experience, coupled with formal graduate and applied training in machine learning and AI governance, I am inspired by BSI’s mission to inspire trust in digital transformation and AI-enabled systems.",
        "Throughout my career, I have specialized in the rigorous verification of complex, mission-critical software systems operating under strict statutory oversight. Supporting the U.S. Department of the Interior (DOI) Acquisition Directorate and IBC for 14 years as a cross-functional Analyst Lead and Senior Escalation Lead, I managed technical system audits, authored standard operating procedures (SOPs), and defended enterprise data integrity through formal federal Inspector General oversight—work recognized with two Department of the Interior Awards for Excellence. This background instilled in me an uncompromising auditing discipline, deep familiarity with traceability matrices, and the ability to review complex technical documentation for regulatory fidelity.",
        "To complement my systems compliance foundation, I have dedicated extensive effort to mastering modern data science, machine learning, and AI regulatory frameworks: (1) Applied Machine Learning & Deep Learning with Spark (UC Berkeley Extension, COMPSCI X433.7 - Grade A), engineering distributed ML pipelines and evaluating neural networks, classifiers, and regression models; (2) Multimodal Architectures & Document Intelligence (Stanford TECH 77), benchmarking Vision Language Models and Document Transformers against OCR baselines with deterministic Pydantic schema validation and prompt-injection anomaly gating; and (3) AI Regulatory Standards & Governance, with formal credentials in ISO/IEC 42001, EU AI Act High-Risk mandates, and NIST AI RMF 1.0 from the AI Tech and Privacy Academy, University of Denver, and Harvard SEAS.",
        "The AI Technical Specialist role requires a unique balance: the technical depth to interrogate complex algorithmic choices, neural network architectures, and training data provenance, alongside the structured communication skills required to articulate non-conformities clearly to client executives and internal assessment teams. Having trained over 85 professionals and led cross-functional technical teams throughout my career, I excel at translating complex technical requirements into defensible, high-integrity audit outcomes.",
        "As an autonomous, self-directed professional operating from the Denver metropolitan area, I welcome the opportunity to support BSI’s global AI Notified Body in shaping the future of AI conformity assessment. Thank you for your consideration, and I look forward to discussing how my experience can contribute to BSI’s distinguished legacy of excellence.",
        "Sincerely,\n\nEbenezer J. Rolston\nDenver, CO | 501-307-3237 | eberolston@yahoo.com"
    ]

    for p_text in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(p_text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(30, 41, 59)
        if p_text.startswith("Dear") or p_text.startswith("Sincerely"):
            run.font.bold = True

    doc.save(docx_path)

def generate_docx_from_md(md_content, docx_path, title_text, subtitle_text):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    r_title = title.add_run(title_text)
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(23, 54, 93)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(10)
    r_sub = sub.add_run(subtitle_text)
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(90, 105, 120)

    lines = md_content.splitlines()

    for line in lines:
        if line.startswith("# "):
            continue
        elif line.startswith("## "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            rh = h.add_run(line[3:])
            rh.font.name = "Calibri"
            rh.font.size = Pt(11.5)
            rh.font.bold = True
            rh.font.color.rgb = RGBColor(23, 54, 93)
        elif line.startswith("### "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
            rh = h.add_run(line[4:])
            rh.font.name = "Calibri"
            rh.font.size = Pt(10.5)
            rh.font.bold = True
            rh.font.color.rgb = RGBColor(40, 70, 110)
        elif line.startswith("* ") or line.startswith("- "):
            b = doc.add_paragraph(style='List Bullet')
            b.paragraph_format.space_before = Pt(0)
            b.paragraph_format.space_after = Pt(2)
            text = line[2:]
            import re
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    rb = b.add_run(part[2:-2])
                    rb.font.bold = True
                else:
                    rb = b.add_run(part)
                rb.font.name = "Calibri"
                rb.font.size = Pt(9.5)
                rb.font.color.rgb = RGBColor(40, 50, 60)
        elif line.startswith("|"):
            continue
        elif line.strip() == "---":
            continue
        elif line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            import re
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    rp = p.add_run(part[2:-2])
                    rp.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    rp = p.add_run(part[1:-1])
                    rp.font.italic = True
                else:
                    rp = p.add_run(part)
                rp.font.name = "Calibri"
                rp.font.size = Pt(9.5)
                rp.font.color.rgb = RGBColor(40, 50, 60)

    doc.save(docx_path)

def main():
    print("Generating BSI AI Technical Specialist Application Package...")

    # 1. Resume
    resume_docx = OUTPUT_DIR / "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Resume.docx"
    resume_md = OUTPUT_DIR / "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Resume.md"
    paragraphs = generate_resume_paragraphs()
    create_docx(resume_docx, paragraphs)
    with open(resume_md, "w") as f:
        f.write(generate_resume_md())
    print(f"  [+] Generated Resume: {resume_docx.name} & {resume_md.name}")

    # 2. Skill Gap Analysis
    sg_docx = OUTPUT_DIR / "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Skill_Gap_Analysis.docx"
    sg_md = OUTPUT_DIR / "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Skill_Gap_Analysis.md"
    sg_text = generate_skill_gap_md()
    with open(sg_md, "w") as f:
        f.write(sg_text)
    generate_docx_from_md(sg_text, sg_docx, "BSI SKILL GAP & FIT ANALYSIS", "AI Technical Specialist • Requisition JR0020733-1 • Candidate: Ebenezer J. Rolston")
    print(f"  [+] Generated Skill Gap Analysis: {sg_docx.name} & {sg_md.name}")

    # 3. Interview Prep
    ip_docx = OUTPUT_DIR / "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Interview_Prep.docx"
    ip_md = OUTPUT_DIR / "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Interview_Prep.md"
    ip_text = generate_interview_prep_md()
    with open(ip_md, "w") as f:
        f.write(ip_text)
    generate_docx_from_md(ip_text, ip_docx, "BSI INTERVIEW PREPARATION PACKAGE", "AI Technical Specialist • Requisition JR0020733-1 • Candidate: Ebenezer J. Rolston")
    print(f"  [+] Generated Interview Prep: {ip_docx.name} & {ip_md.name}")

    # 4. Cover Letter
    cl_docx = OUTPUT_DIR / "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Cover_Letter.docx"
    cl_md = OUTPUT_DIR / "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Cover_Letter.md"
    cl_text = generate_cover_letter_md()
    with open(cl_md, "w") as f:
        f.write(cl_text)
    generate_cover_letter_docx(cl_docx)
    print(f"  [+] Generated Cover Letter: {cl_docx.name} & {cl_md.name}")

    # Synchronize to Twin Directory
    for f in OUTPUT_DIR.iterdir():
        shutil.copy2(f, TWIN_DIR / f.name)
    print(f"  [+] Synchronized all 8 deliverables to twin directory: {TWIN_DIR}")
    print("\nPackage generation completed successfully!")

if __name__ == "__main__":
    main()
