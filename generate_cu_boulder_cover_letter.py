import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    run_name = title.add_run("EBENEZER J. ROLSTON\n")
    run_name.font.name = "Calibri"
    run_name.font.size = Pt(16)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(15, 34, 64) # Navy

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(14)
    run_sub = sub.add_run("Denver, CO • eberolston@yahoo.com • (501) 307-3237 • linkedin.com/in/ebenezer-rolston • github.com/EbeDenver")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(9.5)
    run_sub.font.color.rgb = RGBColor(90, 105, 120)

def generate_cover_letter():
    out_dir = "/Users/eberolston/Downloads/ResumeAgent/output/CU_Boulder_Data_Integration_Analyst"
    os.makedirs(out_dir, exist_ok=True)
    docx_path = os.path.join(out_dir, "Ebenezer_Rolston_CU_Boulder_Data_Integration_Analyst_Cover_Letter.docx")
    md_path = os.path.join(out_dir, "Ebenezer_Rolston_CU_Boulder_Data_Integration_Analyst_Cover_Letter.md")

    # 1. MD Generation
    md_content = """# Ebenezer J. Rolston
**Denver, CO • eberolston@yahoo.com • (501) 307-3237 • [LinkedIn](https://www.linkedin.com/in/ebenezer-rolston/) • [GitHub](https://github.com/EbeDenver)**

October 29, 2026

Search Committee — Data & Integration Analyst (Requisition 75732)  
Office of Information Technology (OIT)  
University of Colorado Boulder  
Boulder, CO 80309  

**Subject: Application for Data & Integration Analyst (Requisition #75732)**

Dear Members of the Search Committee,

I am writing to express my enthusiastic interest in the **Data & Integration Analyst** position (Requisition #75732) within the Office of Information Technology (OIT) at the University of Colorado Boulder. As a Denver-area resident with over two decades of enterprise data integration, ETL/ELT pipeline implementation, and complex systems architecture experience, I am eager to contribute to CU Boulder’s digital transformation. The opportunity to serve as an implementation partner and reliable operational backup to the CRM Data Architect—eliminating single-person dependencies while assuring the integrity of campus-wide CRM integrations—directly aligns with my professional strengths and career commitment to public and educational institutions.

Throughout my career, I have specialized in turning fragmented, heterogeneous data environments into robust, standardized, and auditable pipelines. Most notably, as the **Department ETL Architect for the U.S. Department of the Interior (DOI) FBMS enterprise conversion**, I led the data migration, cleansing, and validation that transitioned legacy bureau databases across four federal agencies into a unified SAP S/4 enterprise platform. In this capacity, I authored comprehensive source-to-target mappings, codified complex business transformation logic, established data lineage documentation, and executed automated reconciliation frameworks to verify byte-level data fidelity. This work was recognized with two Department of the Interior Awards for Excellence.

My technical foundation centers on advanced SQL and relational data architecture across Oracle, PostgreSQL, and SQL Server environments. During my 14-year tenure supporting mission-critical enterprise systems, I served as a Senior Escalation Lead resolving 5 to 10 high-impact technical incidents daily. Rather than applying superficial patches, I conducted rigorous root-cause defect analyses and engineered automated SQL reconciliation scripts that systematically prevented recurring data anomalies. Furthermore, my 11 years at Acxiom Corporation architecting enterprise Customer 360 data pipelines, identity resolution logic, and multi-source data ingestion provided me with deep experience in the foundational integration principles that govern modern CRM and enterprise data ecosystems.

Crucial to the success of CU Boulder’s OIT team is the ability to bridge deep technical architecture with cross-functional campus collaboration. I have extensive experience partnering directly with business analysts, project managers, campus data stewards, and system developers to translate complex institutional requirements into clear, executable technical specifications. Having personally authored standard operating procedures (SOPs) and delivered structured technical training to more than 85 end users and technical staff, I take immense pride in fostering institutional knowledge transfer, documenting workflows, and promoting data quality at the point of origin.

Living in the Denver metropolitan area, I am fully available for regular on-campus collaboration in Boulder as well as hybrid work. I would welcome the opportunity to discuss how my data integration background, advanced SQL capabilities, and dedication to resilient enterprise data architecture can support CU Boulder’s students, faculty, and administration. Thank you for your time, consideration, and leadership.

Sincerely,

**Ebenezer J. Rolston**  
Denver, CO | (501) 307-3237 | eberolston@yahoo.com
"""
    with open(md_path, "w") as f:
        f.write(md_content)

    # 2. DOCX Generation
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    add_header(doc)

    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_after = Pt(10)
    r = p_date.add_run("October 29, 2026\n\nSearch Committee — Data & Integration Analyst (Requisition 75732)\nOffice of Information Technology (OIT)\nUniversity of Colorado Boulder\nBoulder, CO 80309")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(40, 50, 60)

    p_subj = doc.add_paragraph()
    p_subj.paragraph_format.space_after = Pt(10)
    r_subj = p_subj.add_run("Subject: Application for Data & Integration Analyst (Requisition #75732)")
    r_subj.font.name = "Calibri"
    r_subj.font.size = Pt(10.5)
    r_subj.font.bold = True
    r_subj.font.color.rgb = RGBColor(15, 34, 64)

    paragraphs = [
        "Dear Members of the Search Committee,",
        "I am writing to express my enthusiastic interest in the Data & Integration Analyst position (Requisition #75732) within the Office of Information Technology (OIT) at the University of Colorado Boulder. As a Denver-area resident with over two decades of enterprise data integration, ETL/ELT pipeline implementation, and complex systems architecture experience, I am eager to contribute to CU Boulder’s digital transformation. The opportunity to serve as an implementation partner and reliable operational backup to the CRM Data Architect—eliminating single-person dependencies while assuring the integrity of campus-wide CRM integrations—directly aligns with my professional strengths and career commitment to public and educational institutions.",
        "Throughout my career, I have specialized in turning fragmented, heterogeneous data environments into robust, standardized, and auditable pipelines. Most notably, as the Department ETL Architect for the U.S. Department of the Interior (DOI) FBMS enterprise conversion, I led the data migration, cleansing, and validation that transitioned legacy bureau databases across four federal agencies into a unified SAP S/4 enterprise platform. In this capacity, I authored comprehensive source-to-target mappings, codified complex business transformation logic, established data lineage documentation, and executed automated reconciliation frameworks to verify byte-level data fidelity. This work was recognized with two Department of the Interior Awards for Excellence.",
        "My technical foundation centers on advanced SQL and relational data architecture across Oracle, PostgreSQL, and SQL Server environments. During my 14-year tenure supporting mission-critical enterprise systems, I served as a Senior Escalation Lead resolving 5 to 10 high-impact technical incidents daily. Rather than applying superficial patches, I conducted rigorous root-cause defect analyses and engineered automated SQL reconciliation scripts that systematically prevented recurring data anomalies. Furthermore, my 11 years at Acxiom Corporation architecting enterprise Customer 360 data pipelines, identity resolution logic, and multi-source data ingestion provided me with deep experience in the foundational integration principles that govern modern CRM and enterprise data ecosystems.",
        "Crucial to the success of CU Boulder’s OIT team is the ability to bridge deep technical architecture with cross-functional campus collaboration. I have extensive experience partnering directly with business analysts, project managers, campus data stewards, and system developers to translate complex institutional requirements into clear, executable technical specifications. Having personally authored standard operating procedures (SOPs) and delivered structured technical training to more than 85 end users and technical staff, I take immense pride in fostering institutional knowledge transfer, documenting workflows, and promoting data quality at the point of origin.",
        "Living in the Denver metropolitan area, I am fully available for regular on-campus collaboration in Boulder as well as hybrid work. I would welcome the opportunity to discuss how my data integration background, advanced SQL capabilities, and dedication to resilient enterprise data architecture can support CU Boulder’s students, faculty, and administration. Thank you for your time, consideration, and leadership.",
        "Sincerely,\n\nEbenezer J. Rolston\nDenver, CO | (501) 307-3237 | eberolston@yahoo.com"
    ]

    for p_text in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(p_text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(30, 41, 59)
        if p_text.startswith("Dear") or p_text.startswith("Sincerely"):
            run.font.bold = True

    doc.save(docx_path)
    print(f"Generated Cover Letter: {docx_path} & {md_path}")

    # Copy to twin directory
    twin_dir = "/Users/eberolston/Downloads/AI Agent/ResumeAgent/output/CU_Boulder_Data_Integration_Analyst"
    os.makedirs(twin_dir, exist_ok=True)
    import shutil
    shutil.copy2(docx_path, os.path.join(twin_dir, os.path.basename(docx_path)))
    shutil.copy2(md_path, os.path.join(twin_dir, os.path.basename(md_path)))
    print(f"Synchronized Cover Letter to twin directory: {twin_dir}")

if __name__ == "__main__":
    generate_cover_letter()
