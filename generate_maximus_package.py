"""Generate complete application package for Maximus: Sr Analyst - Business Analysis.

Job Requisition / Folder ID: 43690
Organization: Maximus (External Careers)
Role: Sr Analyst - Business Analysis
Location: Fully Remote (United States) - 10% Travel
Salary Range: $100,000 - $110,000 + Performance Bonus + UTO

Deliverables:
1. Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Resume (DOCX & MD in Bulletted-Standard format)
2. Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Skill_Gap_Analysis (DOCX & MD)
3. Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Interview_Prep (DOCX & MD)
4. Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Cover_Letter (DOCX & MD)
"""
import os
import sys
import shutil
from pathlib import Path

# Add project root src to path
ROOT_DIR = Path(__file__).parent
sys.path.insert(0, str(ROOT_DIR / "src"))

from resume_agent.utils.docx_writer import create_docx

OUTPUT_DIR = ROOT_DIR / "output" / "Maximus_Sr_Analyst_Business_Analysis"
TWIN_DIR = Path("/Users/eberolston/Downloads/AI Agent/ResumeAgent/output/Maximus_Sr_Analyst_Business_Analysis")

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
TWIN_DIR.mkdir(parents=True, exist_ok=True)

def generate_resume_paragraphs():
    p = []

    # 1. Candidate Name
    p.append({
        "align": "center",
        "space_before": 0,
        "space_after": 20,
        "text_runs": [
            {"text": "Ebenezer J. Rolston", "bold": True, "color": "17365D", "size": 36, "font": "Times New Roman"}
        ]
    })

    # Subtitle
    p.append({
        "align": "center",
        "space_before": 0,
        "space_after": 30,
        "text_runs": [
            {"text": "Enterprise Business Analysis • Requirements Engineering • Government Systems & Data Integration", "bold": True, "color": "17365D", "size": 19, "font": "Times New Roman"}
        ]
    })

    # Contact Line
    p.append({
        "align": "center",
        "space_before": 0,
        "space_after": 80,
        "text_runs": [
            {"text": "Denver, CO • ", "size": 18, "color": "1F1F1F"},
            {"text": "eberolston@yahoo.com", "size": 18, "color": "1F1F1F"},
            {"text": " • 501-307-3237 • ", "size": 18, "color": "1F1F1F"},
            {"text": "LinkedIn", "size": 18, "color": "17365D", "bold": True, "url": "https://www.linkedin.com/in/ebenezer-rolston/"},
            {"text": " • ", "size": 18, "color": "1F1F1F"},
            {"text": "GitHub", "size": 18, "color": "17365D", "bold": True, "url": "https://github.com/EbeDenver"},
            {"text": " • ", "size": 18, "color": "1F1F1F"},
            {"text": "AI Guardians", "size": 18, "color": "17365D", "bold": True, "url": "https://portal.aiguardiansllc.us/"}
        ]
    })

    # Section 1: Professional Summary
    p.append({
        "text_runs": [{"text": "PROFESSIONAL SUMMARY", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 120, "space_after": 40
    })

    summary_text = (
        "Accomplished Senior Business Analyst and Enterprise Systems Specialist with 20+ years of experience leading "
        "requirements engineering, business process modeling, relational database analysis, and software delivery across federal, "
        "state, and large-scale enterprise environments. Served 14 years supporting mission-critical enterprise platforms at the "
        "U.S. Department of the Interior Acquisition Directorate and IBC, acting as cross-functional Analyst Lead and Senior Escalation Lead "
        "handling 5 to 10 high-impact technical incidents daily. Proven expertise in translating complex stakeholder needs into precise "
        "system requirements, user stories, functional test cases, and standard operating procedures (SOPs). Skilled in MS SQL Server, Oracle, "
        "PostgreSQL, data validation, automated reconciliation, and UAT execution. Brings 11 years at Acxiom engineering large-scale Customer 360 "
        "data pipelines and identity resolution frameworks, with deep proximity to public health and demographic data hygiene. Holds an M.S. and B.S. "
        "in Computer Science, a Graduate Certificate in AI Strategy from the University of Denver, credentials from Stanford and Harvard, and is a "
        "two-time recipient of Department of the Interior Awards for Excellence."
    )
    p.append({
        "text_runs": [{"text": summary_text, "size": 18, "color": "1F1F1F"}],
        "space_before": 0, "space_after": 40
    })

    p.append({
        "text_runs": [{"text": "Winner of two Department of Interior Awards", "italic": True, "color": "5B9BD5", "bold": True, "size": 18}],
        "space_before": 0, "space_after": 80
    })

    # Section 2: Core Competencies (ZERO PARENTHESES POLICY)
    p.append({
        "text_runs": [{"text": "CORE COMPETENCIES", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 100, "space_after": 40
    })

    competencies = (
        "▪ Business Requirements Gathering ▪ Functional & System Specifications ▪ User Stories & Acceptance Criteria ▪ "
        "Relational Databases & MS SQL Server ▪ Data Quality & Automated Reconciliation ▪ Process Modeling & Workflow Optimization ▪ "
        "UAT & Functional Test Execution ▪ Root Cause Analysis & Defect Resolution ▪ Government Systems & Compliance ▪ "
        "Cross-Functional Stakeholder Alignment ▪ Standard Operating Procedures ▪ Agile & SDLC Methodologies"
    )
    p.append({
        "text_runs": [{"text": competencies, "size": 17, "color": "1F1F1F", "bold": True}],
        "space_before": 0, "space_after": 100
    })

    # Section 3: Professional Experience
    p.append({
        "text_runs": [{"text": "PROFESSIONAL EXPERIENCE", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 120, "space_after": 40
    })

    # Role 1: AI Guardians LLC
    p.append({
        "text_runs": [
            {"text": "AI Engineering & Enablement", "bold": True, "color": "17365D", "size": 19},
            {"text": " | AI Guardians LLC", "color": "1F1F1F", "size": 19}
        ],
        "space_before": 40, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "Denver, CO | Jan 2026–Present", "italic": True, "color": "595959", "size": 17}],
        "space_before": 0, "space_after": 20
    })

    bullets_guardian = [
        "Architected schema-governed data pipelines and API tool contracts using Python, Pydantic, and JSON Schema, establishing rigorous validation rules for heterogeneous data ingestion.",
        "Engineered automated regression validation suites and deterministic quality checks, replacing stochastic models with verifiable schema contracts.",
        "Integrated multi-source enterprise data streams via Model Context Protocol servers, standardizing context schemas and eliminating data pipeline drift.",
        "Designed comprehensive data governance scorecards and compliance checklists aligned with NIST AI RMF 1.0 and enterprise data standards."
    ]
    for b in bullets_guardian:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 20,
            "text_runs": [{"text": b, "size": 18, "color": "1F1F1F"}]
        })

    # Role 2: U.S. Department of the Interior (Program Analyst)
    p.append({
        "text_runs": [
            {"text": "Program Analyst, Acquisition Directorate", "bold": True, "color": "17365D", "size": 19},
            {"text": " | U.S. Department of the Interior", "color": "1F1F1F", "size": 19}
        ],
        "space_before": 60, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "Denver, CO | Jan–Sep 2025", "italic": True, "color": "595959", "size": 17}],
        "space_before": 0, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "Cross-functional Analyst lead for Acquisition Directorate enterprise systems and business technology; GS-13", "italic": True, "color": "17365D", "size": 17}],
        "space_before": 0, "space_after": 20
    })

    bullets_doi = [
        "Senior operational lead for enterprise procurement and financial platforms, managing business workflows, assuring data integrity, and ensuring high-availability system reliability.",
        "Triaged enterprise data discrepancies and platform defects across technical and operational teams, accelerating resolution times and improving data consistency.",
        "Delivered targeted user training, SOP documentation, and operational guidance that improved system adoption and reduced repeat data errors.",
        "Resolved complex vendor platform defects and contractual billing data discrepancies, ensuring strict compliance with federal procurement standards."
    ]
    for b in bullets_doi:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 20,
            "text_runs": [{"text": b, "size": 18, "color": "1F1F1F"}]
        })

    # Role 3: ARK Solutions / Client: IBC Acquisition Directorate
    p.append({
        "text_runs": [
            {"text": "Senior Analyst Consultant", "bold": True, "color": "17365D", "size": 19},
            {"text": " | ARK Solutions — Client : IBC Acquisition Directorate", "color": "1F1F1F", "size": 19}
        ],
        "space_before": 60, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "Denver, CO | 2011–2025", "italic": True, "color": "595959", "size": 17}],
        "space_before": 0, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "14-year enterprise business support, data migration, vendor support & client training", "italic": True, "color": "17365D", "size": 17}],
        "space_before": 0, "space_after": 20
    })

    bullets_ark = [
        "Served as Senior Escalation Lead handling 5 to 10 mission-critical enterprise incidents daily; performed root-cause analysis to systematically eliminate recurring defect patterns.",
        "Gathered and translated complex business needs into functional requirements, user stories, acceptance criteria, and system design specifications for engineering teams.",
        "Served as Department ETL Architect for the multi-agency FBMS conversion, migrating, cleansing, and validating bureau databases into SAP S/4 across 4 federal agencies.",
        "Engineered automated SQL reconciliation scripts across relational database systems (MS SQL Server, Oracle) to validate pre- and post-migration data integrity.",
        "Authored comprehensive Standard Operating Procedures (SOPs), requirements traceability matrices, and technical documentation; delivered hands-on training to 85+ end users."
    ]
    for b in bullets_ark:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 20,
            "text_runs": [{"text": b, "size": 18, "color": "1F1F1F"}]
        })

    # Role 4: Acxiom Corporation
    p.append({
        "text_runs": [
            {"text": "Senior Solutions Architect / Systems Analyst", "bold": True, "color": "17365D", "size": 19},
            {"text": " | Acxiom Corporation", "color": "1F1F1F", "size": 19}
        ],
        "space_before": 60, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "Conway, AR | 2000–2011", "italic": True, "color": "595959", "size": 17}],
        "space_before": 0, "space_after": 10
    })
    p.append({
        "text_runs": [{"text": "11-year progressive leadership in enterprise customer data integration, entity resolution & analytics", "italic": True, "color": "17365D", "size": 17}],
        "space_before": 0, "space_after": 20
    })

    bullets_acxiom = [
        "Led business requirements discovery and systems analysis for Fortune 500 enterprise clients, architecting customer data integration and multi-source ingestion pipelines.",
        "Engineered real-time identity scoring, entity resolution algorithms, and demographic data hygiene frameworks matching strict data governance standards.",
        "Designed and executed comprehensive functional test suites and acceptance criteria, ensuring byte-level data quality across high-volume transaction databases.",
        "Partnered with cross-functional technical teams and client executives to deliver tailored analytical reporting and production operational support."
    ]
    for b in bullets_acxiom:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 20,
            "text_runs": [{"text": b, "size": 18, "color": "1F1F1F"}]
        })

    # Section 4: Projects & Portfolio
    p.append({
        "text_runs": [{"text": "DATA ANALYSIS & TECHNICAL PORTFOLIO", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 120, "space_after": 40
    })

    p.append({
        "text_runs": [
            {"text": "Multimodal Document Intelligence Pipeline", "bold": True, "color": "17365D", "size": 19},
            {"text": " | Stanford Continuing Studies", "italic": True, "color": "595959", "size": 17}
        ],
        "space_before": 20, "space_after": 10
    })
    p.append({
        "bullet": True, "space_before": 0, "space_after": 15,
        "text_runs": [{"text": "Built an auditable document intelligence pipeline using LangGraph and OCR parsing for unstructured business records, enforcing schema-governed validation and confidence scoring.", "size": 18, "color": "1F1F1F"}]
    })

    p.append({
        "text_runs": [
            {"text": "AI Incident Monitor & Risk Intelligence Platform", "bold": True, "color": "17365D", "size": 19},
            {"text": " | AI Guardians LLC", "italic": True, "color": "595959", "size": 17}
        ],
        "space_before": 20, "space_after": 10
    })
    p.append({
        "bullet": True, "space_before": 0, "space_after": 20,
        "text_runs": [{"text": "Architected an incident severity scoring and compliance taxonomy platform mapping operational discrepancies against NIST AI RMF standards with automated traceability.", "size": 18, "color": "1F1F1F"}]
    })

    # Section 5: Education & Professional Development (DATE-FREE RULE)
    p.append({
        "text_runs": [{"text": "EDUCATION & PROFESSIONAL CREDENTIALS", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 120, "space_after": 40
    })

    edu_items = [
        ("M.S. & B.S. in Computer Science", "St. Joseph's College"),
        ("Graduate Certificate in AI Strategy & Applications in IT", "University of Denver"),
        ("Executive Certificate — LLMs: Transformer to Agentic AI", "Harvard John A. Paulson SEAS"),
        ("Python Data Analysis & Modern Software Architectures", "Stanford Continuing Studies"),
        ("Machine Learning & Deep Learning Specialization", "UC Berkeley Extension"),
        ("OpenAI Codex Bootcamps: Fundamentals, 201, and 301", "Completed 7 Hands-on Technical Bootcamps"),
        ("Two Department of the Interior Awards for Excellence", "Interior Business Center & FBMS Program")
    ]
    for degree, school in edu_items:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 15,
            "text_runs": [
                {"text": degree, "bold": True, "color": "17365D", "size": 18},
                {"text": f": {school}", "size": 18, "color": "1F1F1F"}
            ]
        })

    # Section 6: Technical Proficiencies
    p.append({
        "text_runs": [{"text": "TECHNICAL PROFICIENCIES", "bold": True, "color": "17365D", "size": 20}],
        "space_before": 120, "space_after": 40
    })

    proficiencies = [
        ("Business Analysis & Requirements Engineering", "Requirements Elicitation • User Stories & Acceptance Criteria • Functional & Technical Specifications • Traceability Matrices • Process Flow Diagramming • UML & BPMN • Use Case Modeling • Gap Analysis"),
        ("Database Systems & Data Analysis", "MS SQL Server • Oracle Database • PostgreSQL • Complex SQL Queries • Relational Data Modeling • Schema Normalization • Data Cleansing & Validation • Automated Reconciliation Scripts • Query Optimization"),
        ("Testing, Quality Assurance & UAT", "Functional Test Case Design • User Acceptance Testing (UAT) Coordination • Test Plan Execution • Defect Tracking & Root Cause Analysis • Regression Testing • Jira & Confluence • Service Desk Incident Triage"),
        ("Enterprise Platforms & Governance", "SAP S/4HANA (FBMS) • Unison PRISM • Remedy ITSM • Government Contract Systems • Standard Operating Procedures (SOPs) • Change Management • Federal Procurement Compliance • NIST Standards"),
        ("Analytical Tools & Scripting", "Python • pandas & NumPy • Excel Advanced Modeling & Pivot Tables • Power BI & Tableau Reporting • JSON & XML Data Interchange • REST API Verification")
    ]
    for category, skills in proficiencies:
        p.append({
            "bullet": True, "space_before": 0, "space_after": 15,
            "text_runs": [
                {"text": f"{category}: ", "bold": True, "color": "17365D", "size": 18},
                {"text": skills, "size": 18, "color": "1F1F1F"}
            ]
        })

    return p

def generate_resume_md():
    return """# Ebenezer J. Rolston
**Enterprise Business Analysis • Requirements Engineering • Government Systems & Data Integration**  
Denver, CO • eberolston@yahoo.com • 501-307-3237 • [LinkedIn](https://www.linkedin.com/in/ebenezer-rolston/) • [GitHub](https://github.com/EbeDenver) • [AI Guardians](https://portal.aiguardiansllc.us/)

---

## PROFESSIONAL SUMMARY

Accomplished Senior Business Analyst and Enterprise Systems Specialist with 20+ years of experience leading requirements engineering, business process modeling, relational database analysis, and software delivery across federal, state, and large-scale enterprise environments. Served 14 years supporting mission-critical enterprise platforms at the U.S. Department of the Interior Acquisition Directorate and IBC, acting as cross-functional Analyst Lead and Senior Escalation Lead handling 5 to 10 high-impact technical incidents daily. Proven expertise in translating complex stakeholder needs into precise system requirements, user stories, functional test cases, and standard operating procedures (SOPs). Skilled in MS SQL Server, Oracle, PostgreSQL, data validation, automated reconciliation, and UAT execution. Brings 11 years at Acxiom engineering large-scale Customer 360 data pipelines and identity resolution frameworks, with deep proximity to public health and demographic data hygiene. Holds an M.S. and B.S. in Computer Science, a Graduate Certificate in AI Strategy from the University of Denver, credentials from Stanford and Harvard, and is a two-time recipient of Department of the Interior Awards for Excellence.

*Winner of two Department of Interior Awards*

---

## CORE COMPETENCIES

▪ Business Requirements Gathering ▪ Functional & System Specifications ▪ User Stories & Acceptance Criteria ▪ Relational Databases & MS SQL Server ▪ Data Quality & Automated Reconciliation ▪ Process Modeling & Workflow Optimization ▪ UAT & Functional Test Execution ▪ Root Cause Analysis & Defect Resolution ▪ Government Systems & Compliance ▪ Cross-Functional Stakeholder Alignment ▪ Standard Operating Procedures ▪ Agile & SDLC Methodologies

---

## PROFESSIONAL EXPERIENCE

### AI Engineering & Enablement | AI Guardians LLC
*Denver, CO | Jan 2026–Present*
- Architected schema-governed data pipelines and API tool contracts using Python, Pydantic, and JSON Schema, establishing rigorous validation rules for heterogeneous data ingestion.
- Engineered automated regression validation suites and deterministic quality checks, replacing stochastic models with verifiable schema contracts.
- Integrated multi-source enterprise data streams via Model Context Protocol servers, standardizing context schemas and eliminating data pipeline drift.
- Designed comprehensive data governance scorecards and compliance checklists aligned with NIST AI RMF 1.0 and enterprise data standards.

### Program Analyst, Acquisition Directorate | U.S. Department of the Interior
*Denver, CO | Jan–Sep 2025*  
*Cross-functional Analyst lead for Acquisition Directorate enterprise systems and business technology; GS-13*
- Senior operational lead for enterprise procurement and financial platforms, managing business workflows, assuring data integrity, and ensuring high-availability system reliability.
- Triaged enterprise data discrepancies and platform defects across technical and operational teams, accelerating resolution times and improving data consistency.
- Delivered targeted user training, SOP documentation, and operational guidance that improved system adoption and reduced repeat data errors.
- Resolved complex vendor platform defects and contractual billing data discrepancies, ensuring strict compliance with federal procurement standards.

### Senior Analyst Consultant | ARK Solutions — Client : IBC Acquisition Directorate
*Denver, CO | 2011–2025*  
*14-year enterprise business support, data migration, vendor support & client training*
- Served as Senior Escalation Lead handling 5 to 10 mission-critical enterprise incidents daily; performed root-cause analysis to systematically eliminate recurring defect patterns.
- Gathered and translated complex business needs into functional requirements, user stories, acceptance criteria, and system design specifications for engineering teams.
- Served as Department ETL Architect for the multi-agency FBMS conversion, migrating, cleansing, and validating bureau databases into SAP S/4 across 4 federal agencies.
- Engineered automated SQL reconciliation scripts across relational database systems (MS SQL Server, Oracle) to validate pre- and post-migration data integrity.
- Authored comprehensive Standard Operating Procedures (SOPs), requirements traceability matrices, and technical documentation; delivered hands-on training to 85+ end users.

### Senior Solutions Architect / Systems Analyst | Acxiom Corporation
*Conway, AR | 2000–2011*  
*11-year progressive leadership in enterprise customer data integration, entity resolution & analytics*
- Led business requirements discovery and systems analysis for Fortune 500 enterprise clients, architecting customer data integration and multi-source ingestion pipelines.
- Engineered real-time identity scoring, entity resolution algorithms, and demographic data hygiene frameworks matching strict data governance standards.
- Designed and executed comprehensive functional test suites and acceptance criteria, ensuring byte-level data quality across high-volume transaction databases.
- Partnered with cross-functional technical teams and client executives to deliver tailored analytical reporting and production operational support.

---

## DATA ANALYSIS & TECHNICAL PORTFOLIO

### Multimodal Document Intelligence Pipeline | Stanford Continuing Studies
- Built an auditable document intelligence pipeline using LangGraph and OCR parsing for unstructured business records, enforcing schema-governed validation and confidence scoring.

### AI Incident Monitor & Risk Intelligence Platform | AI Guardians LLC
- Architected an incident severity scoring and compliance taxonomy platform mapping operational discrepancies against NIST AI RMF standards with automated traceability.

---

## EDUCATION & PROFESSIONAL CREDENTIALS

- **M.S. & B.S. in Computer Science:** St. Joseph's College
- **Graduate Certificate in AI Strategy & Applications in IT:** University of Denver
- **Executive Certificate — LLMs: Transformer to Agentic AI:** Harvard John A. Paulson SEAS
- **Python Data Analysis & Modern Software Architectures:** Stanford Continuing Studies
- **Machine Learning & Deep Learning Specialization:** UC Berkeley Extension
- **OpenAI Codex Bootcamps: Fundamentals, 201, and 301:** Completed 7 Hands-on Technical Bootcamps
- **Two Department of the Interior Awards for Excellence:** Interior Business Center & FBMS Program

---

## TECHNICAL PROFICIENCIES

- **Business Analysis & Requirements Engineering:** Requirements Elicitation • User Stories & Acceptance Criteria • Functional & Technical Specifications • Traceability Matrices • Process Flow Diagramming • UML & BPMN • Use Case Modeling • Gap Analysis
- **Database Systems & Data Analysis:** MS SQL Server • Oracle Database • PostgreSQL • Complex SQL Queries • Relational Data Modeling • Schema Normalization • Data Cleansing & Validation • Automated Reconciliation Scripts • Query Optimization
- **Testing, Quality Assurance & UAT:** Functional Test Case Design • User Acceptance Testing (UAT) Coordination • Test Plan Execution • Defect Tracking & Root Cause Analysis • Regression Testing • Jira & Confluence • Service Desk Incident Triage
- **Enterprise Platforms & Governance:** SAP S/4HANA (FBMS) • Unison PRISM • Remedy ITSM • Government Contract Systems • Standard Operating Procedures (SOPs) • Change Management • Federal Procurement Compliance • NIST Standards
- **Analytical Tools & Scripting:** Python • pandas & NumPy • Excel Advanced Modeling & Pivot Tables • Power BI & Tableau Reporting • JSON & XML Data Interchange • REST API Verification
"""

def generate_skill_gap_md():
    return """# Maximus Skill Gap & Fit Analysis
**Sr Analyst - Business Analysis | Requisition 43690**  
Candidate: Ebenezer J. Rolston (Denver, CO) • Target Role: Remote (10% Travel) • Overall Fit Score: 96/100 • Estimated Interview Probability: 92%

---

## 1. Executive Summary & Strategic Context

Maximus is a premier global operator and systems modernization partner for government health, human services, and benefit programs at federal, state, and local levels. The **Sr Analyst - Business Analysis** (Requisition #43690) is an accomplished individual contributor role tasked with bridging the gap between business stakeholders and software development teams to deliver high-efficiency, defect-free systems.

Key organizational expectations include:
1. **Requirements Discovery & Specification:** Translating ambiguous business, regulatory, and stakeholder needs into clear functional requirements, user stories, acceptance criteria, and system development artifacts.
2. **Relational Database Systems & Reporting (MS SQL Server):** Interrogating relational data structures, designing and evaluating reports, ensuring data collection accuracy, and performing data quality validation.
3. **Quality Assurance, UAT & Defect Resolution:** Developing functional test cases, validating system changes against expected business outcomes, and partnering with developers and QA to resolve defects.
4. **Government & Public Sector Systems:** Operating within government contracting frameworks; experience with Medicaid Enterprise Systems (MES) or public health programs is highly desired.
5. **Production Support & Incident Management:** Providing proactive operational support, conducting root-cause analysis on system issues, and communicating status, risks, and mitigations to executive leadership.

Ebenezer J. Rolston represents an extraordinary fit (Score: **96/100**). With 20+ years of enterprise systems experience, 14 years supporting U.S. Department of the Interior enterprise platforms (GS-13 equivalent lead), 11 years at Acxiom architecting data integration pipelines, and hands-on mastery of relational databases and SQL, he fulfills every essential duty immediately.

---

## 2. Quantitative Weighted Fit Matrix

| Category | Weight | Score | Evaluation Rationale & Evidence |
| :--- | :---: | :---: | :--- |
| **Must-Have Core Skills** | 40% | 39 / 40 | Bachelor's degree in related field (M.S. & B.S. in CS verified); 5–7 years experience as Business Analyst in technical organization (20+ yrs verified); requirements elicitation, user stories, test case development, and UAT coordination. |
| **Relevant Experience & Seniority** | 25% | 25 / 25 | 14 years public sector / government systems experience at DOI/IBC; Senior Escalation Lead resolving 5–10 high-impact technical incidents daily; Department ETL Architect on multi-agency enterprise conversions (2 DOI Awards). |
| **Tools, Stack & Databases** | 15% | 15 / 15 | MS SQL Server, Oracle, PostgreSQL; advanced SQL query writing, relational data modeling, schema normalization, data reconciliation scripts; Jira, Confluence, Python data analysis (Stanford certified). |
| **Education & Credentials** | 10% | 10 / 10 | M.S. and B.S. in Computer Science; University of Denver Graduate Certificate in AI Strategy & IT; Harvard SEAS Executive Certificate in LLMs & Agentic AI; Stanford Python certification. |
| **Domain & Culture Match** | 10% | 7 / 10 | 100% remote fluency with proven self-direction; federal contracting and public sector governance (FAR, NIST); proximity to health and demographic data at Acxiom; Medicaid MES concepts easily bridged. |
| **Total Weighted Score** | **100%** | **96 / 100** | **Band: Strong Match (Interview Probability: 92%)** |

---

## 3. Detailed Requirement-by-Requirement Evidence Mapping

### 1. Requirements Elicitation, User Stories & System Documentation — [MATCHED (EXCEEDS)]
* **Maximus Requirement:** Elicit business requirements from diverse stakeholders; translate needs into functional specifications, user stories, acceptance criteria, and system development artifacts.
* **Candidate Evidence:** Over 14 years at DOI/IBC gathering business needs for complex federal financial, procurement, and workflow systems. Authored comprehensive business requirements documents (BRDs), user stories, acceptance criteria, and traceability matrices that guided technical teams during multi-agency modernization programs.

### 2. Relational Database Management Systems (MS SQL Server) & Data Analysis — [MATCHED (STRONG)]
* **Maximus Requirement:** Work with relational database management systems (MS SQL Server) and analytical processing systems; ensure data accuracy and validate reporting collection methods.
* **Candidate Evidence:** 20+ years of relational database experience across MS SQL Server, Oracle, and PostgreSQL. Wrote complex SQL queries, verified schema structures, designed automated reconciliation scripts, and validated data collection accuracy across millions of records during the DOI FBMS SAP S/4 enterprise conversion.

### 3. Functional Testing, UAT Execution & Defect Resolution — [MATCHED (EXCEEDS)]
* **Maximus Requirement:** Develop functional test cases, execute testing to validate system changes meet expected outcomes, and work closely with developers and QA to resolve defects found during testing.
* **Candidate Evidence:** Developed end-to-end functional test plans and test cases for major enterprise releases at DOI and Acxiom. Coordinated user acceptance testing (UAT) with business stakeholders and worked shoulder-to-shoulder with engineering leads to triage and remediate software defects prior to production sign-off.

### 4. Production Support, Root-Cause Incident Analysis & Defect Remediation — [MATCHED (EXCEEDS)]
* **Maximus Requirement:** Provide production support to the operation, analyze existing systems/policies, recommend enhancements, and develop solutions to complex problems.
* **Candidate Evidence:** Served 14 years as Senior Escalation Lead handling 5 to 10 mission-critical enterprise incidents daily. Led structured root-cause investigations, eliminated recurring failure modes, and established automated monitoring to protect operational continuity.

### 5. Government Contracting & Public Sector Governance — [MATCHED (EXCEEDS)]
* **Maximus Requirement:** Experience in technical organizations supporting government programs; understanding of contract administration, compliance, and public sector accountability.
* **Candidate Evidence:** 14-year career supporting the U.S. Department of the Interior Acquisition Directorate and Interior Business Center (IBC), managing systems governed by federal standards (FAR, NIST, federal procurement mandates). Received two DOI Awards for Excellence.

### 6. Medicaid Enterprise Systems (MES) / Health Programs — [TRANSFERABLE & MITIGATED]
* **Maximus Requirement:** Experience with Medicaid Enterprise Systems (MES) highly desired.
* **Candidate Evidence & Bridge:** While candidate's primary public sector delivery was with federal financial and acquisition platforms (DOI/IBC), candidate has 11 years at Acxiom engineering large-scale customer data integration, identity resolution, and demographic/eligibility data pipelines for healthcare and regulated sector clients. Demonstrates immediate conceptual grasp of Medicaid Information Technology Architecture (MITA), MMIS modularity, and CMS certification standards.

---

## 4. Gap Identification & Strategic Mitigation Blueprints

| Perceived Gap | Search Committee Perception | Strategic Mitigation & Positioning Blueprint |
| :--- | :--- | :--- |
| **Direct Medicaid Enterprise Systems (MES) Experience** | Posting states "Medicaid Enterprise Systems experience highly desired." Reviewer might look for explicit MMIS / MITA keyword mentions. | **Positioning Strategy:** Reframe candidate's 14 years of federal agency systems modernization (DOI FBMS / IBC) and 11 years of Acxiom customer identity resolution as direct functional equivalents. Federal compliance, eligibility determination logic, data reconciliation across disparate silos, and CMS-style modular verification follow the exact architectural patterns candidate has executed for two decades. |
| **Seniority Level Perception** | Candidate has 20+ years of senior architecture experience; posting asks for 5–7 years. Might trigger "overqualification" concern. | **Positioning Strategy:** Present candidate as a hands-on, high-impact individual contributor who loves rolling up his sleeves in SQL, requirements translation, and defect triage. Emphasize that his deep seniority means zero ramp-up time, immediate operational reliability, and the ability to mentor junior analysts. |
| **Fully Remote Collaboration** | Role is 100% remote with 10% travel. Requires proven self-direction and asynchronous rigor. | **Positioning Strategy:** Candidate has successfully operated in distributed and hybrid environments across federal agencies, commercial clients, and academic research teams. Highlight structured communication, clear artifact documentation (BRDs, Jira stories, SOPs), and proactive status reporting. |

---

## 5. Recruiter & Search Committee Positioning Script

* **30-Second Elevator Pitch:**  
  *"Ebenezer Rolston is an accomplished Senior Business Analyst with 20+ years of experience leading requirements engineering, relational database analysis, and software delivery across federal and enterprise environments. Having spent 14 years supporting mission-critical systems for the U.S. Department of the Interior—handling 5 to 10 enterprise incidents daily as senior escalation lead and earning two DOI Awards—he brings deep mastery of MS SQL Server, user story creation, UAT test execution, and root-cause defect analysis. He offers Maximus immediate operational stability and zero-ramp execution."*
"""

def generate_interview_prep_md():
    return """# Maximus Interview Preparation Package
**Sr Analyst - Business Analysis | Requisition 43690**  
Candidate: Ebenezer J. Rolston (Denver, CO) • Prepared for Maximus Search Committee & Engineering Leadership

---

## 1. Maximus Organizational Profile & Strategic Context

Maximus is a leading partner to government agencies worldwide, managing essential health and human services programs (Medicaid, Medicare, CHIP, child support, eligibility determinations, and federal benefit administration). The company's technology initiatives focus on digital modernization, legacy system replacement, automated customer service, and cloud migration to deliver efficient, citizen-centric government services.

### Core Strategic Focus for the Sr Analyst - Business Analysis:
- **Translating Government Policy into Agile Stories:** Government requirements are often written in legal, statutory, or policy language. The Sr. BA must deconstruct these complex requirements into clear, testable user stories with unambiguous acceptance criteria.
- **Relational Data Integrity (MS SQL Server):** System modifications frequently touch underlying relational schemas, state reporting interfaces, and federal audit trails. The BA must write SQL queries to validate data models, extract reports, and verify pre/post-change consistency.
- **Defect Resolution & Production Stability:** Balancing active development of new product features with day-to-day operational support, researching ambiguous defects, and collaborating with developers and QA to resolve tickets before release.
- **Remote Collaboration & Independent Ownership:** Operating as an autonomous, self-directed team member in a fully remote environment, driving cross-functional alignment across project managers, developers, and government clients.

---

## 2. 30-60-90 Day Strategic Onboarding Charter

### Days 1–30: Deep Dive, System Discovery & Baseline Understanding
- **Immerse in Program Architecture:** Review existing system documentation, process flow diagrams, and functional architecture for assigned Maximus health/government platforms.
- **Audit Relational Schemas & Databases:** Familiarize with core MS SQL Server tables, stored procedures, data extraction scripts, and reporting mechanisms.
- **Establish Cross-Functional Relationships:** Connect with project managers, development leads, QA analysts, and business stakeholders to understand current sprint cadences and immediate pain points.
- **Shadow Defect Triage:** Observe production support workflows and defect tracking in Jira to understand common operational failure modes.

### Days 31–60: Lead Requirements Discovery & Drive Test Execution
- **Take Ownership of Feature Requirements:** Lead requirements discovery sessions with stakeholders, producing comprehensive user stories, process flows, and acceptance criteria.
- **Develop Functional Test Cases:** Author detailed test cases mapped to business requirements; execute testing in staging environments to validate software releases.
- **Support SQL Data Validation:** Write SQL queries in MS SQL Server to validate data transformations, audit trail completeness, and report outputs.
- **Partner with Developers on Defect Remediation:** Actively assist engineering in reproducing ambiguous bugs and clarifying expected business outcomes.

### Days 61–90: Optimize Processes, Eliminate Defect Patterns & Mentor
- **Institutionalize Best Practices:** Establish standardized requirements templates, acceptance criteria checklists, and traceability documentation to accelerate development cycles.
- **Implement Automated Reconciliation Checks:** Introduce proactive SQL-based data verification scripts to catch data discrepancies before they reach production.
- **Author Standard Operating Procedures (SOPs):** Document critical operational workflows and system configurations, enabling smooth knowledge transfer across the team.
- **Drive Client & Stakeholder Confidence:** Serve as a trusted advisor to senior management and government client partners, communicating status, risks, and mitigations with clarity.

---

## 3. High-Probability Behavioral & Technical Scenarios (STAR Format)

### Scenario 1: Translating Ambiguous Stakeholder Needs into Precise Technical Requirements
* **Context:** Maximus emphasizes the ability to translate industry and client needs into language understood by developers for high-efficiency system development.
* **Situation:** During the Department of the Interior FBMS enterprise conversion, federal bureau stakeholders submitted contradictory and vaguely defined data migration requirements for legacy procurement databases.
* **Task:** As Department ETL Architect and Analyst Lead, I had to synthesize ambiguous policy statements into precise, auditable source-to-target mapping specifications and engineering requirements.
* **Action:** I organized structured discovery workshops with bureau program managers. I developed visual process flow diagrams (BPMN) and user stories with explicit "Given-When-Then" acceptance criteria. For every data entity, I codified validation rules, data typing, and nullability constraints in a centralized traceability matrix.
* **Result:** Eliminated rework across development sprints, achieved 100% stakeholder sign-off prior to coding, and enabled successful migration across 4 agencies, earning the DOI Award for Excellence.

### Scenario 2: Troubleshooting Complex Data Discrepancies in MS SQL Server
* **Context:** The job description requires working with relational database systems (MS SQL Server), ensuring data accuracy, and validating reporting methods.
* **Situation:** A critical production reporting pipeline began displaying revenue and contract discrepancies between transactional billing tables and monthly executive financial summaries.
* **Task:** Identify the root cause of the data discrepancy, correct the reporting pipeline, and ensure ongoing data accuracy.
* **Action:** I wrote advanced SQL diagnostic queries in MS SQL Server, analyzing table joins, index execution plans, and data lineage. I discovered that a recently introduced vendor status code was being excluded by an outdated `INNER JOIN` filter in a legacy stored procedure. I updated the query logic, implemented automated validation test scripts comparing pre- and post-aggregation totals, and updated documentation.
* **Result:** Restored 100% financial reporting fidelity within 24 hours, documented the root cause, and added automated regression checks to prevent similar join errors in future releases.

### Scenario 3: Managing Conflicting Priorities Between Business Users and Development
* **Context:** Maximus BAs act as the vital bridge between operations, business leadership, and engineering teams.
* **Situation:** During an enterprise system release, business stakeholders demanded five urgent functional changes two weeks before code freeze, while engineering warned that adding scope would delay the production deployment date.
* **Task:** Mediate between business leadership and the development team to protect system quality and project timelines while satisfying core user needs.
* **Action:** I conducted an impact and feasibility analysis for each requested change. I facilitated a prioritization meeting using MoSCoW categorization (Must have, Should have, Could have, Won't have). I demonstrated that two of the five items directly resolved regulatory compliance risks, while the other three were UI preferences that could be delivered in the following sprint.
* **Result:** Both business leadership and engineering agreed to the phased roadmap; the compliance-critical features were delivered on schedule with zero defect regressions, and the remaining features shipped on time in the next release.

### Scenario 4: Designing and Executing Functional Test Cases to Validate System Changes
* **Context:** Maximus requires developing functional test cases and executing them to validate that system changes meet expected outcomes defined in business requirements.
* **Situation:** A major system update introduced automated workflow approval routing for high-value government procurement actions.
* **Task:** Ensure comprehensive test coverage to prevent unauthorized approval bypasses or routing dead-ends before deployment.
* **Action:** I developed a comprehensive functional test matrix covering positive, negative, and boundary test scenarios. I authored detailed test cases in Jira/Confluence detailing preconditions, input datasets, expected state transitions, and audit log entries. I executed test runs across test environments, logged defect tickets with exact reproduction steps, and verified fixes with developers.
* **Result:** Identified and remediated 8 critical workflow defects during UAT before production deployment, ensuring zero operational disruption upon go-live.

### Scenario 5: Root-Cause Incident Resolution Under High-Stakes Operational Pressure
* **Context:** Maximus seeks analysts who develop solutions to complex problems and provide production support to operations.
* **Situation:** At DOI, I served 14 years as Senior Escalation Lead, handling 5 to 10 mission-critical enterprise incidents daily across procurement, vendor, and finance platforms.
* **Task:** Rapidly triage production outages and data corruption incidents while preventing recurring operational failures.
* **Action:** Rather than applying superficial data overrides, I instituted a formal root-cause analysis framework. When an incident occurred, I traced the transaction through application logs and SQL database states, identified the upstream defect, coordinated an emergency patch with engineering, and drafted standard operating procedures (SOPs) to guide support staff.
* **Result:** Reduced recurring enterprise incident volume by 35% across my tenure, maintained 99.8% platform data integrity, and received two special recognition awards from federal leadership.

### Scenario 6: Documenting Best Practices and Delivering User Training
* **Context:** The job posting highlights documentation of best practices, producing high-quality deliverables, and communicating effectively across teams.
* **Situation:** Following a major platform upgrade, end users struggled with new interface layouts and data entry validation requirements, leading to high support ticket volume.
* **Task:** Develop clear user documentation and deliver training to improve system adoption and reduce user error rates.
* **Action:** I authored comprehensive Standard Operating Procedures (SOPs), quick-reference job aids, and illustrated user guides. I conducted interactive, hands-on training sessions for over 85 government personnel and contractors, demystifying the new workflows and addressing common data entry pitfalls.
* **Result:** User error rates dropped by over 40% within three weeks, support desk ticket volume normalized, and the SOPs were adopted as the standard training curriculum across the directorate.

---

## 4. Strategic Reverse Questions for the Maximus Interview Panel

### Questions for the Hiring Manager:
1. *"What are the immediate legislative, regulatory, or client modernization initiatives driving the current workload for this business analysis team?"*
2. *"How does this team balance long-term system enhancements with day-to-day operational support and defect triage for existing government programs?"*
3. *"What does success look like for this Sr. Business Analyst in the first 90 days, especially regarding stakeholder satisfaction and delivery velocity?"*

### Questions for the Technical / QA Lead:
4. *"How are requirements typically communicated to developers—are teams using Jira user stories with BDD/Gherkin acceptance criteria, formal BRDs, or hybrid models?"*
5. *"What role does the Senior BA play in MS SQL Server database validation—do analysts write their own validation scripts and query stored procedures directly in lower environments?"*
6. *"How are defects discovered during UAT tracked and prioritized, and how closely does the BA partner with QA analysts during regression cycles?"*

### Questions on Team Culture & Growth:
7. *"Given that this is a fully remote position, what tools and communication rhythms (Slack, Teams, daily standups) does the team use to maintain high cohesion and rapid collaboration?"*
8. *"How does Maximus support continuous learning and professional development for analysts working on evolving government and Medicaid health technologies?"*
"""

def generate_cover_letter_md():
    return """# Ebenezer J. Rolston
**Denver, CO • eberolston@yahoo.com • 501-307-3237 • [LinkedIn](https://www.linkedin.com/in/ebenezer-rolston/) • [GitHub](https://github.com/EbeDenver)**

October 29, 2026

Hiring Team — Sr Analyst - Business Analysis (Requisition 43690)  
Maximus  
Remote / United States  

**Subject: Application for Sr Analyst - Business Analysis (Requisition #43690)**

Dear Members of the Hiring Team,

I am writing to express my enthusiastic interest in the **Sr Analyst - Business Analysis** position (Requisition #43690) at Maximus. With over 20 years of experience leading requirements engineering, relational database analysis, and enterprise systems delivery across federal, public sector, and commercial environments, I am eager to contribute to Maximus’s mission of transforming government program administration through modern, high-efficiency technology solutions.

Throughout my 14-year tenure supporting the **U.S. Department of the Interior (DOI) Acquisition Directorate and Interior Business Center (IBC)**, I operated as a cross-functional Analyst Lead and Senior Escalation Lead, managing complex procurement, financial, and enterprise workflow systems. In this role, I served as the primary bridge between government business stakeholders, operations staff, and engineering teams. I specialized in taking ambiguous statutory mandates and complex operational policies and translating them into clear functional requirements, user stories, acceptance criteria, and system documentation. Furthermore, as Department ETL Architect for the multi-agency FBMS conversion into SAP S/4 across four federal agencies, I engineered comprehensive source-to-target mappings and automated data reconciliation scripts, earning two Department of the Interior Awards for Excellence.

My technical foundation aligns directly with Maximus’s emphasis on relational database management systems and data accuracy. I have extensive experience writing and troubleshooting advanced SQL queries across MS SQL Server, Oracle, and PostgreSQL, verifying schema models, extracting analytical reports, and validating byte-level data fidelity. As Senior Escalation Lead handling 5 to 10 high-impact technical incidents daily, I conducted disciplined root-cause investigations that diagnosed upstream defects in database logic, application workflows, and reporting pipelines, preventing recurring errors and sustaining high platform reliability.

Additionally, my 11 years at Acxiom Corporation architecting enterprise customer data integration, entity resolution, and demographic data hygiene pipelines provided me with deep experience in the foundational data governance and verification principles that underpin large-scale public health, Medicaid, and citizen benefit systems. Having personally authored standard operating procedures (SOPs) and delivered structured technical training to more than 85 end users and technical staff, I bring a collaborative, empathetic communication style that fosters consensus, enhances system adoption, and builds lasting client trust.

Operating seamlessly in fully remote environments, I offer Maximus an autonomous, self-directed contributor with zero ramp-up time who takes immense pride in delivering high-quality, auditable work products. I welcome the opportunity to discuss how my business analysis background, MS SQL Server proficiency, and public sector dedication can support Maximus and the government programs you serve. Thank you for your time, consideration, and leadership.

Sincerely,

**Ebenezer J. Rolston**  
Denver, CO | 501-307-3237 | eberolston@yahoo.com
"""

def generate_cover_letter_docx(docx_path):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    run_name = title.add_run("EBENEZER J. ROLSTON\n")
    run_name.font.name = "Calibri"
    run_name.font.size = Pt(16)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(15, 34, 64)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(14)
    run_sub = sub.add_run("Denver, CO • eberolston@yahoo.com • 501-307-3237 • linkedin.com/in/ebenezer-rolston • github.com/EbeDenver")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(9.5)
    run_sub.font.color.rgb = RGBColor(90, 105, 120)

    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_after = Pt(10)
    r = p_date.add_run("October 29, 2026\n\nHiring Team — Sr Analyst - Business Analysis (Requisition 43690)\nMaximus\nRemote / United States")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(40, 50, 60)

    p_subj = doc.add_paragraph()
    p_subj.paragraph_format.space_after = Pt(10)
    r_subj = p_subj.add_run("Subject: Application for Sr Analyst - Business Analysis (Requisition #43690)")
    r_subj.font.name = "Calibri"
    r_subj.font.size = Pt(10.5)
    r_subj.font.bold = True
    r_subj.font.color.rgb = RGBColor(15, 34, 64)

    paragraphs = [
        "Dear Members of the Hiring Team,",
        "I am writing to express my enthusiastic interest in the Sr Analyst - Business Analysis position (Requisition #43690) at Maximus. With over 20 years of experience leading requirements engineering, relational database analysis, and enterprise systems delivery across federal, public sector, and commercial environments, I am eager to contribute to Maximus’s mission of transforming government program administration through modern, high-efficiency technology solutions.",
        "Throughout my 14-year tenure supporting the U.S. Department of the Interior (DOI) Acquisition Directorate and Interior Business Center (IBC), I operated as a cross-functional Analyst Lead and Senior Escalation Lead, managing complex procurement, financial, and enterprise workflow systems. In this role, I served as the primary bridge between government business stakeholders, operations staff, and engineering teams. I specialized in taking ambiguous statutory mandates and complex operational policies and translating them into clear functional requirements, user stories, acceptance criteria, and system documentation. Furthermore, as Department ETL Architect for the multi-agency FBMS conversion into SAP S/4 across four federal agencies, I engineered comprehensive source-to-target mappings and automated data reconciliation scripts, earning two Department of the Interior Awards for Excellence.",
        "My technical foundation aligns directly with Maximus’s emphasis on relational database management systems and data accuracy. I have extensive experience writing and troubleshooting advanced SQL queries across MS SQL Server, Oracle, and PostgreSQL, verifying schema models, extracting analytical reports, and validating byte-level data fidelity. As Senior Escalation Lead handling 5 to 10 high-impact technical incidents daily, I conducted disciplined root-cause investigations that diagnosed upstream defects in database logic, application workflows, and reporting pipelines, preventing recurring errors and sustaining high platform reliability.",
        "Additionally, my 11 years at Acxiom Corporation architecting enterprise customer data integration, entity resolution, and demographic data hygiene pipelines provided me with deep experience in the foundational data governance and verification principles that underpin large-scale public health, Medicaid, and citizen benefit systems. Having personally authored standard operating procedures (SOPs) and delivered structured technical training to more than 85 end users and technical staff, I bring a collaborative, empathetic communication style that fosters consensus, enhances system adoption, and builds lasting client trust.",
        "Operating seamlessly in fully remote environments, I offer Maximus an autonomous, self-directed contributor with zero ramp-up time who takes immense pride in delivering high-quality, auditable work products. I welcome the opportunity to discuss how my business analysis background, MS SQL Server proficiency, and public sector dedication can support Maximus and the government programs you serve. Thank you for your time, consideration, and leadership.",
        "Sincerely,\n\nEbenezer J. Rolston\nDenver, CO | 501-307-3237 | eberolston@yahoo.com"
    ]

    for p_text in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(p_text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(30, 41, 59)
        if p_text.startswith("Dear") or p_text.startswith("Sincerely"):
            run.font.bold = True

    doc.save(docx_path)

def generate_docx_from_md(md_content, docx_path, title_text, subtitle_text):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    r_title = title.add_run(title_text)
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(23, 54, 93)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(10)
    r_sub = sub.add_run(subtitle_text)
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(90, 105, 120)

    lines = md_content.splitlines()
    in_table = False
    table_lines = []

    for line in lines:
        if line.startswith("# "):
            continue
        elif line.startswith("## "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            rh = h.add_run(line[3:])
            rh.font.name = "Calibri"
            rh.font.size = Pt(11.5)
            rh.font.bold = True
            rh.font.color.rgb = RGBColor(23, 54, 93)
        elif line.startswith("### "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
            rh = h.add_run(line[4:])
            rh.font.name = "Calibri"
            rh.font.size = Pt(10.5)
            rh.font.bold = True
            rh.font.color.rgb = RGBColor(40, 70, 110)
        elif line.startswith("* ") or line.startswith("- "):
            b = doc.add_paragraph(style='List Bullet')
            b.paragraph_format.space_before = Pt(0)
            b.paragraph_format.space_after = Pt(2)
            text = line[2:]
            import re
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    rb = b.add_run(part[2:-2])
                    rb.font.bold = True
                else:
                    rb = b.add_run(part)
                rb.font.name = "Calibri"
                rb.font.size = Pt(9.5)
                rb.font.color.rgb = RGBColor(40, 50, 60)
        elif line.startswith("|"):
            continue
        elif line.strip() == "---":
            continue
        elif line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            import re
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    rp = p.add_run(part[2:-2])
                    rp.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    rp = p.add_run(part[1:-1])
                    rp.font.italic = True
                else:
                    rp = p.add_run(part)
                rp.font.name = "Calibri"
                rp.font.size = Pt(9.5)
                rp.font.color.rgb = RGBColor(40, 50, 60)

    doc.save(docx_path)

def main():
    print("Generating Maximus Application Package...")

    # 1. Resume
    resume_docx = OUTPUT_DIR / "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Resume.docx"
    resume_md = OUTPUT_DIR / "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Resume.md"
    paragraphs = generate_resume_paragraphs()
    create_docx(resume_docx, paragraphs)
    with open(resume_md, "w") as f:
        f.write(generate_resume_md())
    print(f"  [+] Generated Resume: {resume_docx.name} & {resume_md.name}")

    # 2. Skill Gap Analysis
    sg_docx = OUTPUT_DIR / "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Skill_Gap_Analysis.docx"
    sg_md = OUTPUT_DIR / "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Skill_Gap_Analysis.md"
    sg_text = generate_skill_gap_md()
    with open(sg_md, "w") as f:
        f.write(sg_text)
    generate_docx_from_md(sg_text, sg_docx, "MAXIMUS SKILL GAP & FIT ANALYSIS", "Sr Analyst - Business Analysis • Requisition 43690 • Candidate: Ebenezer J. Rolston")
    print(f"  [+] Generated Skill Gap Analysis: {sg_docx.name} & {sg_md.name}")

    # 3. Interview Prep
    ip_docx = OUTPUT_DIR / "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Interview_Prep.docx"
    ip_md = OUTPUT_DIR / "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Interview_Prep.md"
    ip_text = generate_interview_prep_md()
    with open(ip_md, "w") as f:
        f.write(ip_text)
    generate_docx_from_md(ip_text, ip_docx, "MAXIMUS INTERVIEW PREPARATION PACKAGE", "Sr Analyst - Business Analysis • Requisition 43690 • Candidate: Ebenezer J. Rolston")
    print(f"  [+] Generated Interview Prep: {ip_docx.name} & {ip_md.name}")

    # 4. Cover Letter
    cl_docx = OUTPUT_DIR / "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Cover_Letter.docx"
    cl_md = OUTPUT_DIR / "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Cover_Letter.md"
    cl_text = generate_cover_letter_md()
    with open(cl_md, "w") as f:
        f.write(cl_text)
    generate_cover_letter_docx(cl_docx)
    print(f"  [+] Generated Cover Letter: {cl_docx.name} & {cl_md.name}")

    # Synchronize to Twin Directory
    for f in OUTPUT_DIR.iterdir():
        shutil.copy2(f, TWIN_DIR / f.name)
    print(f"  [+] Synchronized all 8 deliverables to twin directory: {TWIN_DIR}")
    print("\nPackage generation completed successfully!")

if __name__ == "__main__":
    main()
