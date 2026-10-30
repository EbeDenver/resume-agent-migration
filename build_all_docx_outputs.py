"""Build native DOCX outputs for all 4 regenerated packages from their Markdown files.

Packages:
1. BSI AI Technical Specialist
2. Symetra Developer Advocate / AI Enablement
3. CU Boulder Data Integration Analyst
4. Maximus Sr Analyst Business Analysis
"""
import os
import re
import sys
import shutil
from pathlib import Path

# Add project root src to path
ROOT_DIR = Path("/Users/eberolston/Downloads/ResumeAgent")
sys.path.insert(0, str(ROOT_DIR / "src"))

from resume_agent.utils.docx_writer import create_docx
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

TWIN_ROOT = Path("/Users/eberolston/Downloads/AI Agent/ResumeAgent")

def parse_markdown_resume_to_paragraphs(md_text: str):
    """Parses a bulletted-standard markdown resume into OpenXML paragraph dicts."""
    paragraphs = []
    lines = md_text.splitlines()
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
            
        # 1. Candidate Name
        if line.startswith("# "):
            name = line[2:].strip()
            paragraphs.append({
                "align": "center",
                "space_before": 0,
                "space_after": 20,
                "text_runs": [
                    {"text": name, "bold": True, "color": "17365D", "size": 36, "font": "Times New Roman"}
                ]
            })
            i += 1
            continue

        # 2. Subtitle
        if line.startswith("**") and line.endswith("**") and ("•" in line or "|" in line):
            sub = line[2:-2].strip()
            paragraphs.append({
                "align": "center",
                "space_before": 0,
                "space_after": 30,
                "text_runs": [
                    {"text": sub, "bold": True, "color": "17365D", "size": 19, "font": "Times New Roman"}
                ]
            })
            i += 1
            continue

        # 3. Contact Line with Hyperlinks
        if "eberolston@yahoo.com" in line:
            runs = []
            # Parse segments like [LinkedIn](url)
            parts = re.split(r'(\[.*?\]\(.*?\))', line)
            for part in parts:
                link_match = re.match(r'\[(.*?)\]\((.*?)\)', part)
                if link_match:
                    label, url = link_match.groups()
                    runs.append({"text": label, "size": 18, "color": "17365D", "bold": True, "url": url})
                else:
                    if part:
                        runs.append({"text": part, "size": 18, "color": "1F1F1F"})
            paragraphs.append({
                "align": "center",
                "space_before": 0,
                "space_after": 80,
                "text_runs": runs
            })
            i += 1
            continue

        # Horizontal Divider
        if line == "---":
            i += 1
            continue

        # 4. Major Section Header (## )
        if line.startswith("## "):
            sec_title = line[3:].strip()
            paragraphs.append({
                "space_before": 100,
                "space_after": 40,
                "bottom_border": True,
                "text_runs": [
                    {"text": sec_title.upper(), "bold": True, "color": "17365D", "size": 20, "font": "Times New Roman"}
                ]
            })
            i += 1
            continue

        # 5. Award Callout (*Winner of two Department of Interior Awards*)
        if line.startswith("*Winner of two Department of Interior Awards*"):
            paragraphs.append({
                "space_before": 0,
                "space_after": 80,
                "text_runs": [
                    {"text": "Winner of two Department of Interior Awards", "italic": True, "color": "5B9BD5", "bold": True, "size": 18}
                ]
            })
            i += 1
            continue

        # 6. Core Competencies Block (starts with ▪ )
        if line.startswith("▪"):
            paragraphs.append({
                "space_before": 0,
                "space_after": 100,
                "text_runs": [
                    {"text": line, "size": 17, "color": "1F1F1F", "bold": True}
                ]
            })
            i += 1
            continue

        # 7. Subsection Header (### )
        if line.startswith("### "):
            sub_title = line[4:].strip()
            # Check if next line is italic context line
            next_line = lines[i+1].strip() if i+1 < len(lines) else ""
            if next_line.startswith("*") and next_line.endswith("*") and not next_line.startswith("*Winner"):
                context_line = next_line[1:-1].strip()
                paragraphs.append({
                    "space_before": 40,
                    "space_after": 10,
                    "text_runs": [
                        {"text": f"{sub_title} ", "bold": True, "color": "17365D", "size": 18},
                        {"text": context_line, "italic": True, "size": 17, "color": "4A4A4A"}
                    ]
                })
                i += 2
                continue
            else:
                paragraphs.append({
                    "space_before": 40,
                    "space_after": 10,
                    "text_runs": [
                        {"text": sub_title, "bold": True, "color": "17365D", "size": 18}
                    ]
                })
                i += 1
                continue

        # 8. Italic context subtitle on its own
        if line.startswith("*") and line.endswith("*"):
            ctx = line[1:-1].strip()
            paragraphs.append({
                "space_before": 0,
                "space_after": 20,
                "text_runs": [
                    {"text": ctx, "italic": True, "size": 17, "color": "4A4A4A"}
                ]
            })
            i += 1
            continue

        # 9. Bullet Point (- or * )
        if line.startswith("- ") or line.startswith("* "):
            bullet_text = line[2:].strip()
            runs = [{"text": "• "}]
            # Parse bold/italic runs in bullet
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', bullet_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    runs.append({"text": part[2:-2], "bold": True, "size": 18, "color": "1F1F1F"})
                elif part.startswith("*") and part.endswith("*"):
                    runs.append({"text": part[1:-1], "italic": True, "size": 18, "color": "1F1F1F"})
                else:
                    if part:
                        runs.append({"text": part, "size": 18, "color": "1F1F1F"})
            paragraphs.append({
                "space_before": 0,
                "space_after": 20,
                "left_indent": 360,
                "hanging_indent": 180,
                "text_runs": runs
            })
            i += 1
            continue

        # 10. Regular body paragraph (e.g. in Professional Summary)
        runs = []
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', line)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                runs.append({"text": part[2:-2], "bold": True, "size": 18, "color": "1F1F1F"})
            elif part.startswith("*") and part.endswith("*"):
                runs.append({"text": part[1:-1], "italic": True, "size": 18, "color": "1F1F1F"})
            else:
                if part:
                    runs.append({"text": part, "size": 18, "color": "1F1F1F"})
        paragraphs.append({
            "space_before": 0,
            "space_after": 40,
            "text_runs": runs
        })
        i += 1

    return paragraphs


def generate_docx_from_markdown(md_path: Path, docx_path: Path, doc_title: str, doc_subtitle: str):
    """Converts Skill Gap Analysis or Interview Prep Markdown to formatted DOCX."""
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    r_title = title.add_run(doc_title)
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(23, 54, 93)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(10)
    r_sub = sub.add_run(doc_subtitle)
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(90, 105, 120)

    md_content = md_path.read_text(encoding="utf-8")
    lines = md_content.splitlines()

    for line in lines:
        if line.startswith("# "):
            continue
        elif line.startswith("## "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            rh = h.add_run(line[3:])
            rh.font.name = "Calibri"
            rh.font.size = Pt(11.5)
            rh.font.bold = True
            rh.font.color.rgb = RGBColor(23, 54, 93)
        elif line.startswith("### "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
            rh = h.add_run(line[4:])
            rh.font.name = "Calibri"
            rh.font.size = Pt(10.5)
            rh.font.bold = True
            rh.font.color.rgb = RGBColor(40, 70, 110)
        elif line.startswith("* ") or line.startswith("- "):
            b = doc.add_paragraph(style='List Bullet')
            b.paragraph_format.space_before = Pt(0)
            b.paragraph_format.space_after = Pt(2)
            text = line[2:]
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    rb = b.add_run(part[2:-2])
                    rb.font.bold = True
                else:
                    rb = b.add_run(part)
                rb.font.name = "Calibri"
                rb.font.size = Pt(9.5)
                rb.font.color.rgb = RGBColor(40, 50, 60)
        elif line.startswith("|") or line.strip() == "---":
            continue
        elif line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    rp = p.add_run(part[2:-2])
                    rp.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    rp = p.add_run(part[1:-1])
                    rp.font.italic = True
                else:
                    rp = p.add_run(part)
                rp.font.name = "Calibri"
                rp.font.size = Pt(9.5)
                rp.font.color.rgb = RGBColor(40, 50, 60)

    doc.save(docx_path)


def main():
    jobs = [
        {
            "dir_name": "BSI_AI_Technical_Specialist",
            "resume_md": "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Resume.md",
            "resume_docx": "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Resume.docx",
            "gap_md": "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Skill_Gap_Analysis.md",
            "gap_docx": "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Skill_Gap_Analysis.docx",
            "prep_md": "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Interview_Prep.md",
            "prep_docx": "Ebenezer_Rolston_BSI_AI_Technical_Specialist_Interview_Prep.docx",
            "title_gap": "BSI SKILL GAP & FIT ANALYSIS",
            "sub_gap": "AI Technical Specialist • Requisition JR0020733-1 • Candidate: Ebenezer J. Rolston",
            "title_prep": "BSI INTERVIEW PREPARATION PACKAGE",
            "sub_prep": "AI Technical Specialist • Requisition JR0020733-1 • Candidate: Ebenezer J. Rolston"
        },
        {
            "dir_name": "Symetra_Developer_Advocate_AI_Enablement",
            "resume_md": "Ebenezer_Rolston_Symetra_Developer_Advocate_Resume.md",
            "resume_docx": "Ebenezer_Rolston_Symetra_Developer_Advocate_Resume.docx",
            "gap_md": "Ebenezer_Rolston_Symetra_Developer_Advocate_Skill_Gap_Analysis.md",
            "gap_docx": "Ebenezer_Rolston_Symetra_Developer_Advocate_Skill_Gap_Analysis.docx",
            "prep_md": "Ebenezer_Rolston_Symetra_Developer_Advocate_Interview_Prep.md",
            "prep_docx": "Ebenezer_Rolston_Symetra_Developer_Advocate_Interview_Prep.docx",
            "title_gap": "SYMETRA SKILL GAP & FIT ANALYSIS",
            "sub_gap": "Developer Advocate, AI Enablement • Requisition 446718944470 • Candidate: Ebenezer J. Rolston",
            "title_prep": "SYMETRA INTERVIEW PREPARATION PACKAGE",
            "sub_prep": "Developer Advocate, AI Enablement • Requisition 446718944470 • Candidate: Ebenezer J. Rolston"
        },
        {
            "dir_name": "CU_Boulder_Data_Integration_Analyst",
            "resume_md": "Ebenezer_Rolston_CU_Boulder_Data_Integration_Analyst_Resume.md",
            "resume_docx": "Ebenezer_Rolston_CU_Boulder_Data_Integration_Analyst_Resume.docx",
            "gap_md": "Ebenezer_Rolston_CU_Boulder_Data_Integration_Analyst_Skill_Gap_Analysis.md",
            "gap_docx": "Ebenezer_Rolston_CU_Boulder_Data_Integration_Analyst_Skill_Gap_Analysis.docx",
            "prep_md": "Ebenezer_Rolston_CU_Boulder_Data_Integration_Analyst_Interview_Prep.md",
            "prep_docx": "Ebenezer_Rolston_CU_Boulder_Data_Integration_Analyst_Interview_Prep.docx",
            "title_gap": "CU BOULDER SKILL GAP & FIT ANALYSIS",
            "sub_gap": "Data Integration Analyst • Requisition 75732 • Candidate: Ebenezer J. Rolston",
            "title_prep": "CU BOULDER INTERVIEW PREPARATION PACKAGE",
            "sub_prep": "Data Integration Analyst • Requisition 75732 • Candidate: Ebenezer J. Rolston"
        },
        {
            "dir_name": "Maximus_Sr_Analyst_Business_Analysis",
            "resume_md": "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Resume.md",
            "resume_docx": "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Resume.docx",
            "gap_md": "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Skill_Gap_Analysis.md",
            "gap_docx": "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Skill_Gap_Analysis.docx",
            "prep_md": "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Interview_Prep.md",
            "prep_docx": "Ebenezer_Rolston_Maximus_Sr_Analyst_Business_Analysis_Interview_Prep.docx",
            "title_gap": "MAXIMUS SKILL GAP & FIT ANALYSIS",
            "sub_gap": "Sr Analyst - Business Analysis • Requisition 43690 • Candidate: Ebenezer J. Rolston",
            "title_prep": "MAXIMUS INTERVIEW PREPARATION PACKAGE",
            "sub_prep": "Sr Analyst - Business Analysis • Requisition 43690 • Candidate: Ebenezer J. Rolston"
        }
    ]

    for job in jobs:
        out_dir = ROOT_DIR / "output" / job["dir_name"]
        twin_dir = TWIN_ROOT / "output" / job["dir_name"]
        out_dir.mkdir(parents=True, exist_ok=True)
        twin_dir.mkdir(parents=True, exist_ok=True)

        print(f"\nProcessing {job['dir_name']}...")

        # 1. Resume DOCX
        resume_md_path = out_dir / job["resume_md"]
        resume_docx_path = out_dir / job["resume_docx"]
        md_text = resume_md_path.read_text(encoding="utf-8")
        paragraphs = parse_markdown_resume_to_paragraphs(md_text)
        create_docx(resume_docx_path, paragraphs)
        print(f"  [+] Built Resume DOCX: {resume_docx_path.name}")

        # 2. Skill Gap DOCX
        gap_md_path = out_dir / job["gap_md"]
        gap_docx_path = out_dir / job["gap_docx"]
        generate_docx_from_markdown(gap_md_path, gap_docx_path, job["title_gap"], job["sub_gap"])
        print(f"  [+] Built Skill Gap DOCX: {gap_docx_path.name}")

        # 3. Interview Prep DOCX
        prep_md_path = out_dir / job["prep_md"]
        prep_docx_path = out_dir / job["prep_docx"]
        generate_docx_from_markdown(prep_md_path, prep_docx_path, job["title_prep"], job["sub_prep"])
        print(f"  [+] Built Interview Prep DOCX: {prep_docx_path.name}")

        # Copy all updated files to twin directory
        for f in out_dir.iterdir():
            shutil.copy2(f, twin_dir / f.name)
        print(f"  [+] Synchronized files to twin directory: {twin_dir}")

    print("\nAll deliverables regenerated and synchronized successfully!")

if __name__ == "__main__":
    main()
